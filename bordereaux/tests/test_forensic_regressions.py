"""Regression tests for the defects found in the 2026-09-24 forensic
investigation (AI/RESEARCH/TRUEBIND_PRODUCT_REVALIDATION.md sections 3.2
and 4). Each test names the defect ID it locks down. Fixtures are built in
tmp_path, never written into the repository."""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

from bordereaux import dedupe, ingest, mapping, pipeline, schema, validation  # noqa: E402

S = schema


def _canon(rows: list[dict], mapping_: dict[str, str], sheet="s") -> pd.DataFrame:
    raw = pd.DataFrame(rows).astype("string")
    return ingest.apply_mapping(raw, mapping_, sheet_name=sheet)


def _state(mapping_: dict[str, str], sheet="s") -> dict:
    mapped = set(mapping_.values())
    return {sheet: {f.code: ("alias" if f.code in mapped else "unmapped") for f in S.FIELDS}}


def _arith(rows, mapping_):
    c = _canon(rows, mapping_)
    return validation.validate(c, sheet_field_state=_state(mapping_))


# ---------------------------------------------------------------- v5.2 arithmetic (REVALIDATION 3.2)

V52 = {"Ref": S.CLAIM_REF_CODE, "Insured": S.INSURED_NAME_CODE, "Paid this month": S.PAID_MONTH_CODE,
       "Previously paid": S.PREV_PAID_CODE, "Reserve": S.RESERVE_CODE, "Total incurred": S.INCURRED_CODE}


def _v52_row(ref, month, prev, res, inc):
    return {"Ref": ref, "Insured": "A Ltd", "Paid this month": month, "Previously paid": prev,
            "Reserve": res, "Total incurred": inc}


def test_v52_previously_paid_is_included_in_total_incurred():
    """paid this month 100 + previously paid 900 + reserve 500 = 1500 is a MATCH (was a false mismatch)."""
    r = _arith([_v52_row("C1", "100", "900", "500", "1500")], V52)
    assert (r.arithmetic_match_count, r.arithmetic_mismatch_count, r.arithmetic_not_evaluable_count) == (1, 0, 0)


def test_v52_no_previous_payments_and_zero_values():
    r = _arith([_v52_row("C1", "0", "0", "0", "0"), _v52_row("C2", "250", "0", "0", "250")], V52)
    assert (r.arithmetic_match_count, r.arithmetic_mismatch_count) == (2, 0)


def test_v52_legitimate_exception_is_still_flagged():
    r = _arith([_v52_row("C1", "100", "900", "500", "1400")], V52)
    assert r.arithmetic_mismatch_count == 1
    detail = r.exceptions.iloc[0]["detail"]
    assert "previously paid 900.00" in detail and "no fee columns" in detail


def test_this_month_paid_alone_is_never_treated_as_cumulative():
    m = {k: v for k, v in V52.items() if k != "Previously paid"}
    rows = [{k: v for k, v in _v52_row("C1", "100", "900", "500", "1500").items() if k != "Previously paid"}]
    r = _arith(rows, m)
    assert r.arithmetic_not_evaluable_count == 1 and r.arithmetic_mismatch_count == 0
    assert r.not_evaluable_detail.iloc[0]["reason"] == "previously_paid_unmapped"


def test_fees_are_included_when_fee_columns_exist():
    m = {**V52, "Fees this month": S.FEES_PAID_MONTH_CODE, "Fees previously paid": S.FEES_PREV_PAID_CODE,
         "Fee reserve": S.FEES_RESERVE_CODE}
    row = {**_v52_row("C1", "100", "900", "500", "1650"), "Fees this month": "50", "Fees previously paid": "50",
           "Fee reserve": "50"}
    r = _arith([row], m)
    assert (r.arithmetic_match_count, r.arithmetic_mismatch_count) == (1, 0)


def test_partially_mapped_fees_make_total_incurred_not_evaluable():
    m = {**V52, "Fees this month": S.FEES_PAID_MONTH_CODE}
    row = {**_v52_row("C1", "100", "900", "500", "1550"), "Fees this month": "50"}
    r = _arith([row], m)
    assert r.arithmetic_not_evaluable_count == 1
    assert r.not_evaluable_detail.iloc[0]["reason"] == "fees_partially_mapped"


def test_partial_payment_and_missing_values():
    rows = [_v52_row("C1", "", "400", "600", "1000"),   # blank this-month = nil reported
            _v52_row("C2", "100", "0", "", "100"),      # blank reserve = nil reported
            _v52_row("C3", "100", "0", "50", "")]       # blank incurred -> not evaluable
    r = _arith(rows, V52)
    assert (r.arithmetic_match_count, r.arithmetic_not_evaluable_count) == (2, 1)


def test_paid_to_date_column_and_component_consistency():
    m = {**V52, "Paid to date": S.PAID_TD_CODE}
    row = {**_v52_row("C1", "100", "900", "500", "1500"), "Paid to date": "1100"}
    r = _arith([row], m)
    assert r.arithmetic_mismatch_count == 1  # paid to date != this month + previously paid
    assert any("paid to date=1,100.00" in d for d in r.exceptions["detail"])


def test_unparseable_amount_is_never_zero():
    r = _arith([_v52_row("C1", "abc", "900", "500", "1400")], V52)
    assert r.arithmetic_not_evaluable_count == 1 and r.arithmetic_mismatch_count == 0
    assert r.not_evaluable_detail.iloc[0]["reason"] == "paid_this_month_unparseable"


def test_mixed_currency_columns_are_not_added_P9():
    m = {"Ref": S.CLAIM_REF_CODE, "Paid (GBP)": S.PAID_TD_CODE, "Reserve (EUR)": S.RESERVE_CODE,
         "Incurred": S.INCURRED_CODE}
    rows = [{"Ref": "C1", "Paid (GBP)": "100", "Reserve (EUR)": "50", "Incurred": "150"}]
    c = _canon(rows, m)
    assert pd.isna(c.at[0, S.CURRENCY_CODE])
    r = validation.validate(c, sheet_field_state=_state(m))
    assert (r.arithmetic_match_count, r.arithmetic_not_evaluable_count) == (0, 1)
    assert r.not_evaluable_detail.iloc[0]["reason"] == "mixed_currency_columns"


def test_scale_and_currency_transforms_are_recorded():
    m = {"Ref": S.CLAIM_REF_CODE, "Paid (USD m)": S.PAID_TD_CODE}
    c = _canon([{"Ref": "C1", "Paid (USD m)": "1.5"}], m)
    assert c.at[0, S.PAID_TD_CODE] == 1_500_000
    (t,) = c.attrs["transforms"]
    assert (t["currency_from_header"], t["scale_multiplier"], t["source_column"]) == ("USD", 1_000_000.0, "Paid (USD m)")
    assert t["rule_version"]


# ---------------------------------------------------------------- structural rows

def _sheet(tmp_path, rows, name="S", hidden=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = name
    for r in rows:
        ws.append(r)
    if hidden:
        wb.create_sheet("Visible").append(["x"])
        ws.sheet_state = "hidden"
    p = tmp_path / "f.xlsx"
    wb.save(p)
    return ingest.load_workbook_sheets(p)


H7 = ["Claim Ref", "Insured Name", "Date of Loss", "Paid to Date", "Reserve", "Incurred", "Currency"]


def test_total_row_in_narrow_sheet_is_structural_F2(tmp_path):
    rows = [H7] + [[f"C{i}", "Alice", dt.date(2024, 1, 1), 100, 50, 150, "GBP"] for i in range(10)]
    rows += [["Total", None, None, 1000, 500, 1500, None], ["TOTAL CLAIMS", None, None, None, None, 1500, None],
             ["Sum", None, None, 1000, None, None, None]]
    s = _sheet(tmp_path, rows)[0]
    assert len(s.raw) == 10
    assert sorted(er.reason for er in s.excluded_rows) == ["subtotal"] * 3


def test_insured_named_total_is_still_a_claim(tmp_path):
    rows = [H7, ["C1", "Total Logistics Ltd", dt.date(2024, 1, 1), 1, 1, 2, "GBP"]]
    assert len(_sheet(tmp_path, rows)[0].raw) == 1


def test_claim_ref_only_row_is_kept_not_a_title_P1(tmp_path):
    rows = [H7, ["C1", "Alice", dt.date(2024, 1, 1), 1, 1, 2, "GBP"], ["C2"], ["Table B — GBP claims"]]
    s = _sheet(tmp_path, rows)[0]
    assert s.raw["Claim Ref"].tolist() == ["C1", "C2"]
    assert [er.reason for er in s.excluded_rows] == ["title"]


def test_data_after_long_blank_gap_is_not_dropped_P2(tmp_path):
    rows = [H7] + [[f"A{i}", "Alice", dt.date(2024, 1, 1), 1, 1, 2, "GBP"] for i in range(10)]
    rows += [[None] * 7] * 600
    rows += [[f"B{i}", "Bob", dt.date(2024, 1, 1), 1, 1, 2, "GBP"] for i in range(10)]
    s = _sheet(tmp_path, rows)[0]
    assert len(s.raw) == 20
    (run,) = [er for er in s.excluded_rows if er.reason == "blank_run"]
    assert run.count == 600 and run.row_number == 12
    assert s.excluded_row_count == 600


def test_hidden_sheet_is_disclosed_P14(tmp_path):
    s = _sheet(tmp_path, [H7, ["C1", "A", dt.date(2024, 1, 1), 1, 1, 2, "GBP"]], hidden=True)
    hidden = [x for x in s if x.sheet_name == "S"][0]
    assert hidden.hidden and any("hidden" in n for n in hidden.notes)


# ---------------------------------------------------------------- parsing

def test_date_column_is_not_flipped_by_one_iso_value_P5():
    m = {"Ref": S.CLAIM_REF_CODE, "DOL": S.LOSS_DATE_CODE}
    c = _canon([{"Ref": "1", "DOL": "03/04/2024"}, {"Ref": "2", "DOL": "25/06/2024"},
                {"Ref": "3", "DOL": "2024-07-08"}], m)
    assert [d.date().isoformat() for d in c[S.LOSS_DATE_CODE]] == ["2024-04-03", "2024-06-25", "2024-07-08"]


def test_all_ambiguous_date_column_is_disclosed_P5b():
    m = {"Ref": S.CLAIM_REF_CODE, "DOL": S.LOSS_DATE_CODE}
    c = _canon([{"Ref": "1", "DOL": "03/04/2024"}, {"Ref": "2", "DOL": "05/06/2024"}], m)
    assert any("ambiguous" in n for n in c.attrs["parse_notes"])


@pytest.mark.parametrize("values,expected", [
    (["1.234"], [None]),                          # no evidence: ambiguous -> unparseable, never guessed
    (["1.234", "99.50"], [1.234, 99.5]),          # column proves dot-decimal
    (["1.234", "7.500,25"], [1234.0, 7500.25]),   # column proves comma-decimal
    (["1.234.567"], [1234567.0]),
    (["1,234"], [1234.0]),
])
def test_ambiguous_separator_resolved_only_by_column_evidence_P6(values, expected):
    parsed, unparseable = ingest._parse_amount_series(pd.Series(values, dtype="string"))
    got = [None if pd.isna(v) else float(v) for v in parsed]
    assert got == expected
    assert list(unparseable) == [e is None for e in expected]


def test_cp1252_semicolon_csv_with_title_rows_P13(tmp_path):
    p = tmp_path / "x.csv"
    p.write_bytes("Bordereau Q1 2024;;\n;;\nClaim Ref;Insured Name;Paid\nC1;Société Générale;10,5\n".encode("cp1252"))
    (s,) = ingest.load_workbook_sheets(p)
    assert list(s.raw.columns) == ["Claim Ref", "Insured Name", "Paid"]
    assert s.raw.iloc[0]["Insured Name"] == "Société Générale"
    assert any("cp1252" in n for n in s.notes) and any("';'" in n for n in s.notes)


# ---------------------------------------------------------------- mapping

def test_two_columns_for_one_field_is_a_conflict_not_last_wins_P7():
    raw = pd.DataFrame({"Paid": ["100"], "Paid to Date": ["900"]}).astype("string")
    with pytest.raises(ingest.MappingConflictError):
        ingest.apply_mapping(raw, {"Paid": S.PAID_TD_CODE, "Paid to Date": S.PAID_TD_CODE})


def test_fuzzy_conflicts_are_ambiguous_P7b():
    res = mapping.fuzzy_match_headers(["Paid", "Paid to Date", "Amount Paid"])
    bound = [s for s in res.values() if s.field_code == S.PAID_TD_CODE]
    assert len(bound) == 1 and bound[0].review_state == mapping.AMBIGUOUS
    assert all(s.review_state == mapping.AMBIGUOUS for s in res.values())


def test_ambiguous_paid_alias_needs_review():
    (s,) = mapping.fuzzy_match_headers(["Paid"]).values()
    assert s.field_code == S.PAID_TD_CODE and s.review_state == mapping.REVIEW


class _StubAI:
    model = "stub"
    last_usage = None

    def __init__(self, out):
        self.out = out

    def propose(self, headers):
        return self.out


def test_ai_confidence_is_kept_and_conflicts_flagged_P18():
    res = mapping.build_mapping(["Ref Sinistre", "Numéro"],
                                ai_mapper=_StubAI({"Ref Sinistre": (S.CLAIM_REF_CODE, 0.83),
                                                   "Numéro": (S.CLAIM_REF_CODE, 0.61)}))
    by = res.by_column
    assert by["Ref Sinistre"].field_code == S.CLAIM_REF_CODE and by["Ref Sinistre"].confidence == 0.83
    assert by["Ref Sinistre"].review_state == mapping.AMBIGUOUS
    assert by["Numéro"].field_code is None


def test_ai_failure_never_fails_mapping():
    class Boom(_StubAI):
        def propose(self, headers):
            raise RuntimeError("provider down")
    res = mapping.build_mapping(["Mystery column"], ai_mapper=Boom({}))
    assert res.suggestions[0].field_code is None and "provider down" in res.ai_unavailable_reason


def test_ai_output_is_schema_validated():
    out = mapping._validate_ai_output(
        {"mappings": [{"source_column": "A", "field_code": "CR0104M", "confidence": 0.9},
                      {"source_column": "B", "field_code": "DROP TABLE", "confidence": 0.9},
                      {"source_column": "not requested", "field_code": "CR0104M", "confidence": 0.9},
                      {"source_column": "C", "field_code": "CR0104M", "confidence": 7}]}, {"A", "B", "C"})
    assert out == {"A": ("CR0104M", 0.9)}


# ---------------------------------------------------------------- periods, duplicates, sheet classes

def _run(path):
    sheets = pipeline.load_workbook(path)
    props = pipeline.propose_mapping_for_workbook(sheets)
    conf = {p.sheet.sheet_name: {s.source_column: s.field_code for s in p.mapping.suggestions if s.field_code}
            for p in props}
    return pipeline.run_workbook_pipeline(sheets, conf, props)


_NAMES = ["Acorn", "Birchwood", "Cobalt", "Driftwood", "Ember", "Foxglove", "Granite", "Heron", "Ivory", "Juniper",
          "Kestrel", "Larch", "Meridian", "Nimbus", "Obsidian", "Pinnacle", "Quarry", "Rowan", "Saffron", "Tundra"]


def test_multi_period_workbook_is_not_duplicates_F1(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for mth in range(1, 4):
        ws = wb.create_sheet(f"2024-{mth:02d}")
        ws.append(H7)
        for i in range(20):
            ws.append([f"CLM{i}", f"{_NAMES[i]} Holdings Ltd", dt.date(2023, 5, 1 + i), 100 * mth, 50,
                       100 * mth + 50, "GBP"])
    p = tmp_path / "m.xlsx"
    wb.save(p)
    r = _run(p)
    assert r.duplicates.empty
    assert r.health.exact_duplicates == 0


def test_same_period_repeat_is_exact_and_unknown_period_is_review(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sender A"
    ws.append(H7)
    ws.append(["C1", "A Ltd", dt.date(2024, 1, 1), 1, 1, 2, "GBP"])
    ws.append(["C1", "A Ltd", dt.date(2024, 1, 1), 1, 1, 2, "GBP"])
    ws2 = wb.create_sheet("Sender B")
    ws2.append(H7)
    ws2.append(["C1", "A Ltd", dt.date(2024, 1, 1), 1, 1, 2, "GBP"])
    p = tmp_path / "d.xlsx"
    wb.save(p)
    kinds = sorted(_run(p).duplicates["match_type"])
    assert kinds == ["exact_duplicate", "repeat_period_unknown"]


def test_period_from_text():
    assert dedupe.period_from_text("Mar 2024") == "2024-03"
    assert dedupe.period_from_text("Bordereau 202403") == "2024-03"
    assert dedupe.period_from_text("Sender A") is None


def test_partially_recognised_foreign_sheet_is_not_excluded_F3(tmp_path):
    wb = openpyxl.Workbook()
    en = wb.active
    en.title = "EN"
    en.append(H7)
    for i in range(10):
        en.append([f"E{i}", "Alice", dt.date(2024, 1, 1), 1, 1, 2, "GBP"])
    fr = wb.create_sheet("FR")
    fr.append(["Référence sinistre", "Nom de l'assuré", "Date du sinistre", "Payé à ce jour", "Réserve",
               "Total encouru", "Devise"])
    for i in range(10):
        fr.append([f"F{i}", "Bob", dt.date(2024, 1, 1), 1, 1, 2, "EUR"])
    p = tmp_path / "ml.xlsx"
    wb.save(p)
    r = _run(p)
    statuses = {a.sheet_name: a.status for a in r.coverage.sheet_audit}
    assert statuses["FR"] == "partial"
    assert len(r.canonical) == 20  # the French rows stay in the claim set
    assert r.coverage.reconciliation.non_claim_summary_rows == 0
    assert not r.health.score_reliable


def test_limits_reject_instead_of_truncating(tmp_path):
    rows = [H7] + [[f"C{i}", "A", dt.date(2024, 1, 1), 1, 1, 2, "GBP"] for i in range(50)]
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    p = tmp_path / "big.xlsx"
    wb.save(p)
    with pytest.raises(ingest.WorkbookLimitError):
        ingest.load_workbook_sheets(p, limits=ingest.ReadLimits(max_rows_total=20))
    assert len(ingest.load_workbook_sheets(p)[0].raw) == 50
