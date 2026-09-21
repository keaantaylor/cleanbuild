"""Regression test for the "parenthetical suffix breaks alias matching"
defect: a header like "Paid to Date (GBP)" scored just under
FUZZY_THRESHOLD against its own alias ("paid to date") because the
"(GBP)" suffix dragged the token_sort_ratio down, so real monetary
columns went unmapped and every arithmetic check on that sheet came back
"not evaluable" -- reported as a data problem when it was an ingestion
defect. Covers a full sheet (every monetary column suffixed, not one
header in isolation) plus the currency-hint population when the sheet
has no separate Currency column."""

from __future__ import annotations

import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

import openpyxl  # noqa: E402

from bordereaux import pipeline  # noqa: E402
from bordereaux.mapping import fuzzy_match_headers  # noqa: E402

FIXTURE = REPO_ROOT / "tests" / "fixtures" / "suffixed_headers.xlsx"
FIXTURE_EUR = REPO_ROOT / "tests" / "fixtures" / "suffixed_headers_eur.xlsx"
FIXTURE_CONFLICT = REPO_ROOT / "tests" / "fixtures" / "suffixed_headers_conflict.xlsx"


def _build_fixture() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Claim Ref", "Insured Name", "Paid to Date (GBP)", "Reserve Amount (GBP)", "Incurred Amount (GBP)"])
    ws.append(["C1", "Alice Ltd", 100.0, 50.0, 150.0])
    ws.append(["C2", "Bob Co", 200.0, 0.0, 200.0])
    ws.append(["C3", "Carl Inc", 300.0, 25.0, 325.0])
    wb.save(FIXTURE)


def _build_eur_fixture() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Claim Ref", "Insured Name", "Paid Amount (EUR)", "Reserve Amount (EUR)",
               "Incurred Amount (EUR)", "Currency (EUR)"])
    ws.append(["C1", "Alice Ltd", 100.0, 50.0, 150.0, "EUR"])
    ws.append(["C2", "Bob Co", 200.0, 0.0, 200.0, "EUR"])
    wb.save(FIXTURE_EUR)


def test_eur_suffixed_headers_map_and_parse() -> None:
    """EUR variant of the GBP fixture above, plus an explicit 'Currency
    (EUR)' header -- exercises the direct currency-column mapping path,
    not just the currency-hint-from-amount-header fallback."""
    _build_eur_fixture()
    sheets = pipeline.load_workbook(FIXTURE_EUR)
    proposals = pipeline.propose_mapping_for_workbook(sheets)
    proposal = proposals[0]

    mapped = {s.source_column: s.field_code for s in proposal.mapping.suggestions if s.field_code}
    assert len(mapped) == 6, f"every EUR-suffixed header should still alias-match, got {mapped}"
    assert mapped["Currency (EUR)"] == "CR0110CM", "an explicit 'Currency (EUR)' header must map to the currency field"
    print(f"OK: all 6 EUR-suffixed headers mapped: {mapped}")

    confirmed = {"Sheet1": mapped}
    result = pipeline.run_workbook_pipeline(sheets, confirmed, proposals, source_name=FIXTURE_EUR.name)
    canonical = result.canonical
    assert canonical["CR0126CM"].notna().all(), "Paid amounts must all parse despite the (EUR) suffix"
    assert canonical["CR0130CM"].notna().all(), "Reserve amounts must all parse despite the (EUR) suffix"
    assert canonical["CR0155CM"].notna().all(), "Incurred amounts must all parse despite the (EUR) suffix"
    assert (canonical["CR0110CM"] == "EUR").all(), (
        f"currency should read EUR from the explicit 'Currency (EUR)' column, got {canonical['CR0110CM'].tolist()}"
    )
    print("OK: EUR-suffixed monetary columns + explicit currency column all parsed correctly")


def _build_conflict_fixture() -> None:
    """Two sheets in one workbook, one entirely in GBP, one entirely in
    EUR, neither with a separate Currency column -- the currency-hint
    inference (from the amount header's own suffix) must fire correctly
    and independently per sheet, never bleeding one sheet's currency
    into the other's."""
    wb = openpyxl.Workbook()
    gbp = wb.active
    gbp.title = "UK_Book"
    gbp.append(["Claim Ref", "Insured Name", "Paid Amount (GBP)", "Reserve Amount (GBP)", "Incurred Amount (GBP)"])
    gbp.append(["G1", "Alice Ltd", 100.0, 50.0, 150.0])
    gbp.append(["G2", "Bob Co", 200.0, 0.0, 200.0])

    eur = wb.create_sheet("EU_Book")
    eur.append(["Claim Ref", "Insured Name", "Paid Amount (EUR)", "Reserve Amount (EUR)", "Incurred Amount (EUR)"])
    eur.append(["E1", "Carl SARL", 300.0, 75.0, 375.0])
    eur.append(["E2", "Dana SA", 400.0, 0.0, 400.0])
    wb.save(FIXTURE_CONFLICT)


def test_gbp_and_eur_sheets_in_one_workbook_do_not_conflict() -> None:
    """The currency-conflict scenario, rerun now that suffix mapping
    actually works -- a result recorded against the earlier, broken
    mapping would have been invalid, since every suffixed header came
    back unmapped and the currency-hint inference path this test
    exercises could never have fired at all."""
    _build_conflict_fixture()
    sheets = pipeline.load_workbook(FIXTURE_CONFLICT)
    proposals = pipeline.propose_mapping_for_workbook(sheets)
    confirmed = {
        p.sheet.sheet_name: {s.source_column: s.field_code for s in p.mapping.suggestions if s.field_code}
        for p in proposals
    }
    for p in proposals:
        got = confirmed[p.sheet.sheet_name]
        assert len(got) == 5, f"{p.sheet.sheet_name}: every suffixed header should alias-match, got {got}"

    result = pipeline.run_workbook_pipeline(sheets, confirmed, proposals, source_name=FIXTURE_CONFLICT.name)
    canonical = result.canonical
    by_sheet_currency = canonical.groupby("_source_sheet")["CR0110CM"].unique().apply(list).to_dict()
    assert by_sheet_currency["UK_Book"] == ["GBP"], f"UK_Book must read GBP only, got {by_sheet_currency['UK_Book']}"
    assert by_sheet_currency["EU_Book"] == ["EUR"], f"EU_Book must read EUR only, got {by_sheet_currency['EU_Book']}"

    exceptions = result.validation_result.exceptions
    inconsistency = exceptions[exceptions["rule"] == "currency_inconsistency"] if not exceptions.empty else exceptions
    assert inconsistency.empty, (
        f"GBP-sheet and EUR-sheet claims must never be treated as the same claim reporting two currencies, "
        f"got {inconsistency.to_dict('records')}"
    )
    print("OK: GBP-suffixed and EUR-suffixed sheets in one workbook resolve independently, no cross-contamination")


def test_currency_suffix_casing_whitespace_and_punctuation_variants() -> None:
    """Forensic-report edge cases: lowercase, mixed casing, extra
    whitespace, underscores in place of spaces, and no space before the
    parenthetical must all still resolve to the same canonical field --
    normalize_header casefolds and collapses separators before matching,
    so none of these should behave differently from the plain case."""
    cases = {
        "Paid Amount (GBP)": "CR0126CM",
        "paid amount (gbp)": "CR0126CM",
        "PAID AMOUNT (Gbp)": "CR0126CM",
        "Paid_Amount_(GBP)": "CR0126CM",
        "Paid Amount(EUR)": "CR0126CM",
        "  Paid Amount   (GBP)  ": "CR0126CM",
        "Paid-Amount-(GBP)": "CR0126CM",
        "Paid Amount ( GBP )": "CR0126CM",
        "Reserve Amount (EUR)": "CR0130CM",
        "Incurred Amount (EUR)": "CR0155CM",
        "Currency (EUR)": "CR0110CM",
    }
    for header, expected_code in cases.items():
        got = fuzzy_match_headers([header])[header]
        assert got.field_code == expected_code, (
            f"{header!r} should map to {expected_code}, got {got.field_code} (method={got.method})"
        )
    print(f"OK: {len(cases)} currency-suffix casing/whitespace/punctuation variants all map correctly")


def main() -> None:
    _build_fixture()
    sheets = pipeline.load_workbook(FIXTURE)
    proposals = pipeline.propose_mapping_for_workbook(sheets)
    proposal = proposals[0]

    mapped = {s.source_column: s.field_code for s in proposal.mapping.suggestions if s.field_code}
    assert len(mapped) == 5, f"every suffixed header should still alias-match, got {mapped}"
    print(f"OK: all 5 suffixed headers mapped: {mapped}")

    confirmed = {"Sheet1": mapped}
    result = pipeline.run_workbook_pipeline(sheets, confirmed, proposals, source_name=FIXTURE.name)

    canonical = result.canonical
    assert canonical["CR0126CM"].notna().all(), "Paid amounts must all parse despite the (GBP) suffix"
    assert canonical["CR0130CM"].notna().all(), "Reserve amounts must all parse despite the (GBP) suffix"
    assert canonical["CR0155CM"].notna().all(), "Incurred amounts must all parse despite the (GBP) suffix"
    assert (canonical["CR0110CM"] == "GBP").all(), (
        f"currency should be populated from the header hint since no Currency column exists, "
        f"got {canonical['CR0110CM'].tolist()}"
    )
    print("OK: all monetary columns parsed correctly, currency populated from the header hint")

    assert result.health.arithmetic_not_evaluable == 0, (
        f"no row should be not-evaluable -- every input mapped and parsed cleanly, "
        f"got {result.health.arithmetic_not_evaluable}"
    )
    assert result.health.arithmetic_mismatches == 0, "paid+reserve==incurred on every row of this fixture"
    print("OK: arithmetic reconciliation runs cleanly -- previously this whole sheet came back not-evaluable")

    test_currency_suffix_casing_whitespace_and_punctuation_variants()
    test_eur_suffixed_headers_map_and_parse()
    test_gbp_and_eur_sheets_in_one_workbook_do_not_conflict()

    print("\nHeader-suffix regression test PASSED.")


if __name__ == "__main__":
    main()
