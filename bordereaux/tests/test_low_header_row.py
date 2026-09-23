"""TB-002 regression: a sheet whose header sits well below row 5 -- e.g.
beneath a title band and a logo occupying the sheet's first several rows
-- must still be found and processed, not silently dropped as if the
sheet never existed. Previously HEADER_SCAN_ROWS = 5 meant a header at
row 7 (0-indexed 6) was never even considered."""

from __future__ import annotations

import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

import openpyxl  # noqa: E402

from bordereaux import ingest  # noqa: E402

FIXTURE_DIR = REPO_ROOT / "tests" / "fixtures"


def _build_low_header_fixture() -> Path:
    path = FIXTURE_DIR / "low_header_row.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logo_Image_Header"
    # Rows 1-6: title band a logo/letterhead would occupy above the real
    # header -- sparse, non-tabular text, exactly what a banner looks
    # like to the row-scoring function.
    ws.append(["ACME REINSURANCE GROUP"])
    ws.append(["Bordereau — Q3 2026"])
    ws.append([])
    ws.append(["Prepared by: Claims Ops"])
    ws.append([])
    ws.append([])
    # Row 7 (0-indexed 6): the real header.
    ws.append(["Claim Ref", "Insured Name", "Paid Amount", "Reserve Amount", "Incurred Amount"])
    for i in range(300):
        ws.append([f"CLM-IMG-{i + 1:05d}", f"Insured {i}", 100.0 + i, 50.0, 150.0 + i])
    wb.save(path)
    return path


def test_header_seven_rows_down_is_found_and_all_rows_survive() -> None:
    path = _build_low_header_fixture()
    sheets = ingest.load_workbook_sheets(path)
    assert len(sheets) == 1, "the sheet must still be enumerated even before header detection runs"

    sheet = sheets[0]
    assert not sheet.skipped, f"sheet was skipped: {sheet.skip_reason!r}"
    assert sheet.header_row_index == 6, f"expected header at 0-indexed row 6, got {sheet.header_row_index}"
    assert len(sheet.raw) == 300, f"expected 300 claim rows, got {len(sheet.raw)}"

    refs = set(sheet.raw["Claim Ref"])
    expected = {f"CLM-IMG-{i + 1:05d}" for i in range(300)}
    assert refs == expected, "every CLM-IMG-* row must survive intact"
    print(f"OK: header found at row {sheet.header_row_index + 1}; all {len(sheet.raw)} rows present")


def _build_deep_preamble_fixture() -> Path:
    """Reproduces a real Lloyd's-style bordereau preamble: several
    paragraphs of syndicate/broker/coverholder detail pushing the header
    to row 41 (0-indexed 40) -- past the 30-row window this file's
    HEADER_SCAN_ROWS was raised to before this fix, still not generous
    enough."""
    path = FIXTURE_DIR / "deep_preamble_header_row_41.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "05_Header_Row_41"
    for i in range(40):  # rows 1-40: preamble
        ws.append([f"Lloyd's Syndicate Bordereau -- preamble line {i + 1}"])
    ws.append(["Claim Ref", "Insured Name", "Paid Amount", "Reserve Amount", "Incurred Amount"])
    for i in range(120):
        ws.append([f"CLM-{i + 1:05d}", f"Insured {i}", 100.0 + i, 50.0, 150.0 + i])
    wb.save(path)
    return path


def test_header_forty_one_rows_down_is_found_and_all_rows_survive() -> None:
    path = _build_deep_preamble_fixture()
    sheets = ingest.load_workbook_sheets(path)
    assert len(sheets) == 1

    sheet = sheets[0]
    assert not sheet.skipped, f"sheet was skipped: {sheet.skip_reason!r}"
    assert sheet.header_row_index == 40, f"expected header at 0-indexed row 40, got {sheet.header_row_index}"
    assert len(sheet.raw) == 120, f"expected 120 claim rows, got {len(sheet.raw)}"
    print(f"OK: header found at row {sheet.header_row_index + 1}; all {len(sheet.raw)} rows present")


if __name__ == "__main__":
    test_header_seven_rows_down_is_found_and_all_rows_survive()
    test_header_forty_one_rows_down_is_found_and_all_rows_survive()
    print("\nLow-header-row regression tests PASSED.")
