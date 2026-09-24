"""Regression tests for Section 5/9/10 of the forensic repair brief:
row-count reconciliation and per-sheet audit classification. Reproduces
the reported shape (a large opaque-header sheet, a large foreign-
language sheet, a normal sheet with a duplicate and a rejected row) and
asserts every number the reconciliation is supposed to answer, plus
that raw data for a fully-unmapped sheet survives into the export."""

from __future__ import annotations

import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

import openpyxl  # noqa: E402

from bordereaux import pipeline  # noqa: E402

import tempfile  # noqa: E402

_GEN_DIR = Path(tempfile.gettempdir()) / "truebind_generated_fixtures"  # never rewrite tracked files
_GEN_DIR.mkdir(parents=True, exist_ok=True)
FIXTURE = _GEN_DIR / "reconciliation_scenario.xlsx"


def _build_fixture() -> Path:
    wb = openpyxl.Workbook()

    gibberish = wb.active
    gibberish.title = "Unmappable_Gibberish"
    gibberish.append(["ZX_001", "ZX_002", "ZX_003", "ZX_004", "ZX_005", "ZX_006", "ZX_007", "ZX_008"])
    for i in range(200):
        gibberish.append([f"G{i:05d}", f"N{i}", 100.0 + i, 50.0, 150.0 + i, "2024-01-01", "open", "GBP"])

    fr = wb.create_sheet("FR_Cedante_Full")
    fr.append(["Référence sinistre", "Nom de l'assuré", "Référence police", "Date de survenance",
               "Montant réglé", "Montant provision", "Montant total encouru", "Devise"])
    for i in range(900):
        fr.append([f"FR{i:06d}", f"Client {i} SARL", f"POL-{i:05d}", "2024-02-01",
                   100.0 + i, 50.0, 150.0 + i, "EUR"])

    normal = wb.create_sheet("Normal")
    normal.append(["Claim Ref", "Insured Name", "Paid Amount", "Reserve Amount", "Incurred Amount"])
    normal.append(["C1", "Alice Ltd", 100.0, 50.0, 150.0])
    normal.append(["C1", "Alice Ltd", 100.0, 50.0, 150.0])  # exact duplicate (same claim ref)
    normal.append(["", "", None, None, None])  # blank row -- rejected before mapping

    wb.save(FIXTURE)
    return FIXTURE


def _run():
    _build_fixture()
    sheets = pipeline.load_workbook(FIXTURE)
    proposals = pipeline.propose_mapping_for_workbook(sheets)
    confirmed = {
        p.sheet.sheet_name: {s.source_column: s.field_code for s in p.mapping.suggestions if s.field_code}
        for p in proposals
    }
    # Simulate a reviewer who hasn't (yet) manually mapped either forensic sheet.
    confirmed["Unmappable_Gibberish"] = {}
    confirmed["FR_Cedante_Full"] = {}
    return pipeline.run_workbook_pipeline(sheets, confirmed, proposals, source_name=FIXTURE.name), sheets


def test_sheet_audit_classifies_every_sheet_correctly() -> None:
    result, _ = _run()
    by_name = {rec.sheet_name: rec for rec in result.coverage.sheet_audit}

    gib = by_name["Unmappable_Gibberish"]
    assert gib.status == "unmapped"
    assert gib.fields_mapped == 0
    assert gib.rows_processed == 200
    assert gib.rows_rejected == 0
    assert "no recognized business fields" in gib.reason

    fr = by_name["FR_Cedante_Full"]
    assert fr.status == "unmapped"
    assert fr.rows_processed == 900

    normal = by_name["Normal"]
    assert normal.status == "mapped", normal.reason
    assert normal.rows_processed == 2, "2 real claim rows (the exact-dupe pair), the blank row rejected"
    assert normal.rows_rejected == 1
    print("OK: sheet audit correctly classifies mapped/unmapped sheets with accurate row/field counts")


def test_reconciliation_totals_match_and_reconcile() -> None:
    result, _ = _run()
    recon = result.coverage.reconciliation
    assert recon is not None

    assert recon.source_worksheets == 3
    assert recon.mapped_rows == 2
    assert recon.unmapped_rows == 1100, "200 opaque-header + 900 French-header rows"
    assert recon.rejected_rows == 1, "the one blank row on the Normal sheet"
    assert recon.duplicate_rows == 2, "the exact-duplicate pair"
    assert recon.exported_rows == 1102, "every processed row reaches canonical, mapped or not"
    assert recon.source_data_rows == 1103, "1100 unmapped + 2 mapped + 1 rejected"
    assert recon.rows_requiring_review == 1102, "1100 unmapped-sheet rows + 2 duplicate rows"
    assert recon.reconciles, f"reconciliation should hold exactly, discrepancy={recon.discrepancy}"
    print(f"OK: reconciliation totals match and reconcile exactly -- {recon.as_lines()}")


def test_unmapped_sheet_raw_data_survives_into_export(tmp_path=None) -> None:
    """Section 12: a fully-unmapped sheet's ORIGINAL headers and values
    must be recoverable from the export -- not just its row count."""
    import openpyxl as _openpyxl

    result, sheets = _run()
    out_dir = _GEN_DIR / "_recon_export_tmp"
    paths = pipeline.write_workbook_outputs(result, out_dir, "reconciliation_scenario", sheets=sheets)

    wb = _openpyxl.load_workbook(paths["health"])
    assert "Unmapped-Unmappable_Gibberish" in wb.sheetnames
    assert "Unmapped-FR_Cedante_Full" in wb.sheetnames

    gib_ws = wb["Unmapped-Unmappable_Gibberish"]
    assert [c.value for c in gib_ws[1]] == ["ZX_001", "ZX_002", "ZX_003", "ZX_004",
                                              "ZX_005", "ZX_006", "ZX_007", "ZX_008"]
    assert gib_ws.max_row - 1 == 200, "all 200 raw rows must be recoverable, not just the row count"
    assert gib_ws[2][0].value == "G00000", "actual received values must be recoverable, not just headers"

    assert "Sheet audit" in wb.sheetnames
    import shutil
    shutil.rmtree(out_dir, ignore_errors=True)
    print("OK: fully-unmapped sheets' original headers and values are recoverable from the export")


if __name__ == "__main__":
    test_sheet_audit_classifies_every_sheet_correctly()
    test_reconciliation_totals_match_and_reconcile()
    test_unmapped_sheet_raw_data_survives_into_export()
    print("\nReconciliation/audit-trail regression tests PASSED.")
