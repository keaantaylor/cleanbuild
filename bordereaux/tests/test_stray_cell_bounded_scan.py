"""TB-005 regression: a single stray value far outside a sheet's real
data must never inflate the read into an unbounded scan of Excel's
absolute row/column limits. Before this fix, one cell at A1048576
turned a ~20-row sheet into a 1,048,576-row read; this test asserts
both that it completes quickly and that the reported row count reflects
the sheet's true data extent, not the inflated declared range."""

from __future__ import annotations

import sys
import time
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

import openpyxl  # noqa: E402

from bordereaux import ingest  # noqa: E402

# Generated fixtures go to a temp dir: tests must never rewrite tracked files.
import tempfile  # noqa: E402

FIXTURE_DIR = Path(tempfile.gettempdir()) / "truebind_generated_fixtures"
FIXTURE_DIR.mkdir(parents=True, exist_ok=True)


def _build_stray_cell_fixture() -> Path:
    path = FIXTURE_DIR / "stray_cell.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stray_Last_Row"
    ws.append(["Claim Ref", "Insured Name", "Paid Amount", "Reserve Amount", "Incurred Amount"])
    for i in range(20):
        ws.append([f"CLM-{i + 1:05d}", f"Insured {i}", 100.0 + i, 50.0, 150.0 + i])
    # A single value at Excel's absolute last row -- inflates the sheet's
    # declared used-range to the full 1,048,576 rows with no real data
    # anywhere in between.
    ws["A1048576"] = "stray"
    wb.save(path)
    return path


def test_stray_cell_does_not_trigger_unbounded_scan() -> None:
    path = _build_stray_cell_fixture()

    started = time.monotonic()
    sheets = ingest.load_workbook_sheets(path)
    elapsed = time.monotonic() - started

    assert elapsed < 5.0, f"bounded read should complete in a few seconds, took {elapsed:.1f}s"
    assert len(sheets) == 1
    sheet = sheets[0]
    assert not sheet.skipped, f"sheet was skipped: {sheet.skip_reason!r}"
    assert len(sheet.raw) == 20, f"expected 20 real claim rows, got {len(sheet.raw)}"
    # The stray cell is real sheet content: since the P2 fix (never stop
    # reading silently) it is read and accounted for -- the 1M-row blank gap
    # is ONE collapsed ledger entry (not a million stored rows) and the stray
    # value is excluded with a stated reason, never dropped unseen.
    runs = [er for er in sheet.excluded_rows if er.reason == "blank_run"]
    assert len(runs) == 1 and runs[0].count > 1_000_000, runs
    stray = [er for er in sheet.excluded_rows if er.values.get("Claim Ref") == "stray"]
    assert len(stray) == 1 and stray[0].row_number == 1_048_576, stray
    assert len(sheet.excluded_rows) < 10, "blank rows must be collapsed, not stored one by one"
    print(f"OK: bounded scan completed in {elapsed:.2f}s; {len(sheet.raw)} real rows; stray cell at row "
          f"{stray[0].row_number} accounted for ({stray[0].reason}); {runs[0].count} blank rows collapsed")


if __name__ == "__main__":
    test_stray_cell_does_not_trigger_unbounded_scan()
    print("\nStray-cell bounded-scan regression test PASSED.")
