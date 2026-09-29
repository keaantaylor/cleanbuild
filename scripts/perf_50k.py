#!/usr/bin/env python3
"""P1.7 performance gate: a 50,000-row workbook, end to end, through the real
API and worker, while the API keeps answering.

    python scripts/perf_50k.py --database-url postgresql+psycopg://... [--port 8799]

Starts uvicorn (embedded worker, each job in its own child process) against
the given database, signs up, uploads a seeded 50k-row workbook, confirms the
proposed mapping, processes it, and meanwhile samples GET /health/ready every
250 ms. Prints one JSON line and exits non-zero unless:
  - upload -> report COMPLETE takes < 120 s, and
  - no health probe during the run took >= 1 s (the API was never blocked).
"""

from __future__ import annotations

import argparse
import datetime
import json
import os
import random
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

import httpx
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
BACKEND = ROOT / "truebind-web" / "backend"
ROWS = 50_000
BUDGET_S = 120.0
MAX_PROBE_S = 1.0
HEADER = [
    "Claim Reference", "Policy Number", "Insured Name", "Date of Loss", "Date Notified", "Claim Status", "Currency",
    "Paid to Date", "Outstanding Reserve", "Total Incurred", "Loss Location",
]  # fmt: skip
# Realistic shape for a delegated-authority claims bordereau: company-style
# insured names (a few large repeat clients, a long tail of small ones), loss
# dates across three underwriting years, notification after the loss.
_FIRST = [
    "Northfield",
    "Bluewater",
    "Harrington",
    "Silverline",
    "Oakridge",
    "Meridian",
    "Castlebay",
    "Redstone",
    "Fairview",
    "Ashworth",
    "Kingsley",
    "Brightmoor",
    "Thornbury",
    "Westgate",
    "Ironwood",
    "Sandpiper",
    "Greenhaven",
    "Lockwood",
    "Marlowe",
    "Pemberton",
    "Rosewood",
    "Stonebridge",
    "Underwood",
    "Whitfield",
    "Yardley",
    "Abbotsford",
    "Beaumont",
    "Carrington",
    "Dunmore",
    "Elmstead",
    "Foxhall",
    "Glenville",
    "Hawthorne",
    "Inglewood",
    "Juniper",
    "Kestrel",
    "Langford",
    "Millbrook",
    "Newhaven",
    "Orchard",
    "Penrose",
    "Quayside",
    "Riverside",
    "Summit",
    "Trafalgar",
    "Uplands",
    "Valeview",
    "Waverley",
    "Zenith",
    "Albion",
    "Bramley",
    "Clifton",
    "Drummond",
    "Easton",
    "Fenwick",
    "Galway",
    "Holloway",
    "Ivybridge",
    "Jarrow",
    "Kilbride",
]
_SECOND = ["Shipping", "Logistics", "Freight", "Trading", "Holdings", "Marine", "Transport", "Industries", "Foods",
           "Engineering", "Construction", "Retail", "Hotels", "Farms", "Motors", "Textiles", "Energy", "Estates",
           "Healthcare", "Media"]  # fmt: skip
_SUFFIX = ["Ltd", "Group", "Inc", "PLC", "Co", "LLP"]
_CITIES = ["London", "Dublin", "Manchester", "Cork", "Glasgow", "Rotterdam", "Hamburg", "Lyon", "Madrid", "Milan"]


def make_adversarial_workbook(path: Path) -> None:
    """Every insured name alike ("Insured N") and every loss date in one
    month: ~10^8 probable-duplicate candidate pairs. Must still complete
    in budget, with that check reported as not assessed."""
    rng = random.Random(7)  # noqa: S311 -- deterministic test data, not security
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Claims")
    ws.append(HEADER)
    for i in range(ROWS):
        paid, res = round(rng.uniform(0, 50_000), 2), round(rng.uniform(0, 20_000), 2)
        day = i % 28 + 1
        ws.append([
            f"CLM-{i:06d}", f"POL-{i % 997:05d}", f"Insured {i % 997}", f"2024-03-{day:02d}", f"2024-04-{day:02d}",
            rng.choice(["Open", "Closed"]), rng.choice(["GBP", "EUR", "USD"]), paid, res, round(paid + res, 2),
            "London",
        ])  # fmt: skip
    wb.save(path)


def make_workbook(path: Path) -> None:
    rng = random.Random(7)  # noqa: S311 -- deterministic test data, not security
    insureds = [f"{f} {s} {x}" for f in _FIRST for s in _SECOND for x in _SUFFIX]
    rng.shuffle(insureds)
    weights = [1.0 / (k + 1) ** 0.8 for k in range(len(insureds))]  # a few big repeat clients, a long tail
    names = rng.choices(insureds, weights=weights, k=ROWS)
    policy_no = {n: k for k, n in enumerate(insureds)}
    start = datetime.date(2022, 1, 1)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Claims")
    ws.append(HEADER)
    for i, name in enumerate(names):
        paid, res = round(rng.uniform(0, 50_000), 2), round(rng.uniform(0, 20_000), 2)
        loss = start + datetime.timedelta(days=rng.randrange(3 * 365))
        notified = loss + datetime.timedelta(days=rng.randrange(1, 90))
        ws.append([
            f"CLM-{i:06d}", f"POL-{policy_no[name]:05d}-{loss.year}", name, loss.isoformat(),
            notified.isoformat(), rng.choice(["Open", "Closed", "Reopened"]), rng.choice(["GBP", "EUR", "USD"]),
            paid, res, round(paid + res, 2), rng.choice(_CITIES),
        ])  # fmt: skip
    wb.save(path)


class Prober(threading.Thread):
    def __init__(self, base: str) -> None:
        super().__init__(daemon=True)
        self.base, self.stop = base, threading.Event()
        self.samples: list[float] = []

    def run(self) -> None:
        with httpx.Client(timeout=10) as c:
            while not self.stop.is_set():
                t0 = time.monotonic()
                try:
                    c.get(f"{self.base}/health/ready")
                    self.samples.append(time.monotonic() - t0)
                except httpx.HTTPError:
                    self.samples.append(10.0)
                self.stop.wait(0.25)


def wait_status(c: httpx.Client, rid: str, wanted: str, deadline: float) -> dict[str, object]:
    r: dict[str, object] = {}
    while time.monotonic() < deadline:
        r = c.get(f"/api/v1/reports/{rid}").json()
        if r["status"] == wanted:
            return dict(r)
        if r["status"] in ("FAILED", "CANCELLED"):
            raise RuntimeError(f"report {r['status']}: {r.get('processing_error')}")
        time.sleep(0.5)
    raise TimeoutError(f"report did not reach {wanted}: {json.dumps(r, default=str)[:2000]}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--database-url", required=True)
    ap.add_argument("--port", type=int, default=8799)
    ap.add_argument("--shape", choices=("realistic", "adversarial"), default="realistic")
    ap.add_argument("--redis-url", default="", help="enable Redis job wake-ups for the server under test")
    args = ap.parse_args()
    work = Path(tempfile.mkdtemp(prefix="truebind-perf-"))
    book = work / "claims_50k.xlsx"
    (make_adversarial_workbook if args.shape == "adversarial" else make_workbook)(book)
    env = {
        **os.environ, "DATABASE_URL": args.database_url, "TRUEBIND_ENV": "test", "TRUEBIND_NO_DOTENV": "1",
        "TRUEBIND_EMBEDDED_WORKER": "1", "ALLOW_SIGNUP": "1", "TRUEBIND_DATA_DIR": str(work / "data"),
        "TRUEBIND_STORAGE_DIR": str(work / "objects"), "REDIS_URL": args.redis_url,
    }  # fmt: skip
    log = open(work / "server.log", "wb")  # noqa: SIM115 -- closed in finally
    result: dict[str, object] = {"rows": ROWS, "shape": args.shape, "server_log": str(work / "server.log")}
    subprocess.run([sys.executable, "-m", "alembic", "upgrade", "head"], cwd=BACKEND, env=env, check=True)
    server = subprocess.Popen(  # noqa: S603
        [sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(args.port)],
        cwd=BACKEND, env=env, stdout=log, stderr=subprocess.STDOUT,
    )  # fmt: skip
    base = f"http://127.0.0.1:{args.port}"
    try:
        for _ in range(120):
            try:
                if httpx.get(f"{base}/health/ready", timeout=2).status_code == 200:
                    break
            except httpx.HTTPError:
                pass
            time.sleep(0.5)
        with httpx.Client(base_url=base, timeout=300) as c:
            me = c.post("/api/v1/auth/signup", json={
                "email": f"perf-{int(time.time())}@example.com", "password": "correct horse battery staple",
                "display_name": "Perf", "organisation": "Perf Org"}).json()  # fmt: skip
            c.headers["X-CSRF-Token"] = me["csrf_token"]
            prober = Prober(base)
            prober.start()
            t0 = time.monotonic()
            with open(book, "rb") as f:
                up = c.post("/api/v1/reports/upload", files={"file": ("claims_50k.xlsx", f)})
            up.raise_for_status()
            rid = up.json()["id"]
            deadline = t0 + BUDGET_S * 2
            wait_status(c, rid, "WAITING_FOR_REVIEW", deadline)
            t_ingest = time.monotonic() - t0
            for s in c.get(f"/api/v1/reports/{rid}/sheets").json():
                if s["status"] == "SKIPPED":
                    continue
                m = c.get(f"/api/v1/reports/{rid}/sheets/{s['id']}/mapping").json()
                choices = {fld["field_code"]: fld["source_column"] for fld in m["fields"]}
                c.post(f"/api/v1/reports/{rid}/sheets/{s['id']}/mapping", json={"mappings": choices}).raise_for_status()
            c.post(f"/api/v1/reports/{rid}/process").raise_for_status()
            report = wait_status(c, rid, "COMPLETE", deadline)
            total = time.monotonic() - t0
            summary = c.get(f"/api/v1/reports/{rid}/summary").json()["summary"]
            result["not_assessed_checks"] = [ch["check"] for ch in summary.get("not_assessed_checks", [])]
            prober.stop.set()
            prober.join()
            probes = sorted(prober.samples)
            result.update({
                "ingest_s": round(t_ingest, 1), "total_s": round(total, 1),
                "rows_processed": report["rows_processed"], "probes": len(probes),
                "probe_max_s": round(probes[-1], 3) if probes else None,
                "probe_p95_s": round(probes[int(len(probes) * 0.95) - 1], 3) if probes else None,
            })  # fmt: skip
    finally:
        server.terminate()
        server.wait(timeout=30)
        log.close()
    # A hostile file must finish too -- by declaring the check it could not
    # afford as not assessed, never by silently skipping it.
    expected_not_assessed = ["probable_duplicates"] if args.shape == "adversarial" else []
    ok = (
        result.get("rows_processed") == ROWS
        and result.get("not_assessed_checks") == expected_not_assessed
        and float(str(result["total_s"])) < BUDGET_S
        and float(str(result["probe_max_s"])) < MAX_PROBE_S
    )
    result["passed"] = ok
    print(json.dumps(result))  # noqa: T201 -- CLI output
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
