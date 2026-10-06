"""Open a report's workbook in a connected provider and read edits back.

- Open: the working copy (the original plus every approved correction) is
  uploaded as a NEW file; the source file is never overwritten. The bytes it
  started from are kept (by hash) so edits can be compared later.
- Pull: the remote file is downloaded and compared cell by cell with those
  bytes on the confirmed claims sheets. Every changed cell becomes a PROPOSED
  correction that goes through the execution policy like any other: nothing
  edited outside TrueBind is applied without the approval it needs, and
  blocked changes (formulas, header rows, claim references) are refused.
- Push: the latest working copy is written over the remote file as a new
  version (the provider keeps its history).
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from .. import connectors
from ..models._util import utcnow
from ..models.connectors import ConnectorLink
from ..models.corrections import Correction
from ..models.reports import Report, Sheet
from . import audit_service, corrections_service, deliverables
from .storage import derived_key, get_store

MAX_PULL_CHANGES = 2000


class ConnectorUnavailable(RuntimeError):
    pass


def _provider(key: str) -> connectors.Connector:
    p = connectors.get(key)
    if p is None:
        raise ConnectorUnavailable("Unknown provider.")
    if not p.configured():
        raise ConnectorUnavailable(f"{p.label} is not connected on this server.")
    return p


def working_copy(db: Session, report: Report) -> bytes:
    """The original plus every approved correction (or the original alone)."""
    try:
        body, _ = corrections_service.corrected_bytes(db, report)
        return body
    except corrections_service.CorrectionError:
        wb, _ = deliverables.load_original(db, report)
        buf = io.BytesIO()
        wb.save(buf)
        return corrections_service._deterministic(buf.getvalue())


def _key(report: Report, sha: str) -> str:
    return derived_key(report.tenant_id, report.source_sha256 or "0" * 64, f"connector-base-{sha[:16]}")


def open_in(db: Session, report: Report, provider: str, actor: str, actor_user_id: str | None) -> ConnectorLink:
    p = _provider(provider)
    body = working_copy(db, report)
    sha = hashlib.sha256(body).hexdigest()
    name = f"{Path(report.file_name).stem} (TrueBind {dt.date.today().isoformat()}).xlsx"
    remote = p.upload(name, body)
    get_store().put_derived(_key(report, sha), body, db=db)
    link = ConnectorLink(tenant_id=report.tenant_id, report_id=report.id, provider=p.key, file_id=remote.file_id,
                         web_url=remote.web_url, base_sha256=sha, created_by=actor)
    db.add(link)
    db.flush()
    audit_service.log_action(db, report.tenant_id, report.id, "CONNECTOR_EXPORTED", "CONNECTOR_LINK", link.id,
                             after={"provider": p.key, "file_id": remote.file_id, "base_sha256": sha},
                             actor=actor, actor_user_id=actor_user_id)
    return link


def push(db: Session, report: Report, link: ConnectorLink, actor: str, actor_user_id: str | None) -> ConnectorLink:
    p = _provider(link.provider)
    body = working_copy(db, report)
    sha = hashlib.sha256(body).hexdigest()
    p.upload(Path(report.file_name).stem + ".xlsx", body, file_id=link.file_id)
    get_store().put_derived(_key(report, sha), body, db=db)
    link.base_sha256 = sha
    audit_service.log_action(db, report.tenant_id, report.id, "CONNECTOR_PUSHED", "CONNECTOR_LINK", link.id,
                             after={"provider": p.key, "file_id": link.file_id, "base_sha256": sha},
                             actor=actor, actor_user_id=actor_user_id)
    return link


def diff(base: bytes, edited: bytes, sheets: dict[str, int]) -> list[tuple[str, str, str | None, str | None]]:
    """(sheet, cell, before, after) for every value that differs on the given
    sheets. Edits on or above the header row come back too; the execution
    policy refuses them (BLOCKED), so they are reported, never applied."""
    a = openpyxl.load_workbook(io.BytesIO(base))
    b = openpyxl.load_workbook(io.BytesIO(edited))
    out = []
    for name in sheets:
        if name not in a.sheetnames or name not in b.sheetnames:
            continue
        wa, wb = a[name], b[name]
        rows, cols = max(wa.max_row, wb.max_row), max(wa.max_column, wb.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                va, vb = wa.cell(r, c).value, wb.cell(r, c).value
                if _norm(va) != _norm(vb):
                    out.append((name, wa.cell(r, c).coordinate, _text(va), _text(vb)))
                    if len(out) > MAX_PULL_CHANGES:
                        return out
    return out


def _text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _norm(v):
    t = _text(v)
    return None if t is None or t.strip() == "" else t.strip()


def pull(db: Session, report: Report, link: ConnectorLink, actor: str, actor_user_id: str | None) -> dict:
    p = _provider(link.provider)
    base = get_store().get_derived(_key(report, link.base_sha256), db=db)
    if base is None:
        raise ConnectorUnavailable("The copy this file started from is no longer available; open it again.")
    edited = p.download(link.file_id)
    sheets = {s.sheet_name: (s.header_row_index or 0) + 1
              for s in db.query(Sheet).filter(Sheet.report_id == report.id, Sheet.status == "CONFIRMED")}
    changes = diff(base, edited, sheets)
    if len(changes) > MAX_PULL_CHANGES:
        raise ConnectorUnavailable(f"More than {MAX_PULL_CHANGES} cells changed; send the file as a resubmission.")
    pending = {(c.sheet_name, c.cell) for c in db.query(Correction).filter(Correction.report_id == report.id,
                                                                           Correction.status == "PROPOSED")}
    proposed, blocked, skipped = [], [], 0
    for sheet, cell, _before, after in changes:
        if (sheet, cell) in pending:
            skipped += 1
            continue
        try:
            c = corrections_service.propose(db, report, sheet, cell, after,
                                            f"Edited in {p.label} by a reviewer; read back by {actor}", actor,
                                            source="connector")
            proposed.append({"id": c.id, "sheet": sheet, "cell": cell, "after": after, "policy": c.policy})
        except corrections_service.CorrectionError as exc:
            blocked.append({"sheet": sheet, "cell": cell, "reason": str(exc)})
    link.last_pulled_at = utcnow()
    summary = {"changed_cells": len(changes), "proposed": len(proposed), "blocked": len(blocked),
               "already_pending": skipped}
    audit_service.log_action(db, report.tenant_id, report.id, "CONNECTOR_PULLED", "CONNECTOR_LINK", link.id,
                             after={"provider": p.key, **summary}, actor=actor, actor_user_id=actor_user_id)
    return {**summary, "items": proposed[:200], "blocked_items": blocked[:200]}


def link_out(link: ConnectorLink) -> dict:
    p = connectors.get(link.provider)
    return {"id": link.id, "provider": link.provider, "label": p.label if p else link.provider,
            "web_url": link.web_url, "base_sha256": link.base_sha256, "created_by": link.created_by,
            "created_at": link.created_at.isoformat(),
            "last_pulled_at": link.last_pulled_at.isoformat() if link.last_pulled_at else None}
