"""Customer deliverables built from a processed report and its stored original:

- the annotated workbook: the customer's own workbook, layout kept, every
  claim-sheet cell coloured by result, a Review notes column, hover comments
  with the fix, an Issues sheet grouped by type with links to each cell, and a
  Review Summary sheet first (COUNTIFS counts and a legend). Source values are
  never altered.
- the corrected copy ("mend"): only safe, reversible fixes (dates held as text
  or serial numbers where unambiguous, currency variants to ISO codes, trimmed
  spaces, zero-padded policy numbers, status synonyms, amounts held as text),
  each listed on a Change Log sheet. Duplicates, arithmetic, reserves and
  anything ambiguous are never touched.
- the query letter to the sender: the sender-owned findings grouped by issue,
  with row references.
- the month-on-month check against an earlier report.

Text that starts with = + - @ is always written as text, never as a formula.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections import defaultdict
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from sqlalchemy.orm import Session

from bordereaux.domain_config import CLAIM_STATUS_SYNONYMS
from bordereaux.rules import SEVERITY_ORDER
from bordereaux.rules import rule as catalogue_rule
from bordereaux.schema import FIELDS_BY_CODE

from ..models.reports import ClaimRow, ExcludedRow, Mapping, Report, Sheet, ValidationResult
from .storage import IntegrityError, get_store

GREEN, AMBER, RED, GREY = "E2F0D9", "FFF2CC", "FFC7CE", "E7E6E6"
RED_TEXT = "9C0006"
FILL = {k: PatternFill("solid", start_color=v, end_color=v) for k, v in
        {"ok": GREEN, "warn": AMBER, "err": RED, "grey": GREY}.items()}
DARK_RED = Font(color=RED_TEXT)
BOLD = Font(bold=True)
MAX_COMMENTS = 10_000
RESULT = {"FAIL": "Error", "REVIEW": "Warning", "NOT_EVALUABLE": "Couldn't check"}
TONE_OF = {"FAIL": "err", "REVIEW": "warn", "NOT_EVALUABLE": "warn"}  # can't validate shows amber in the workbook
_FORMULA = ("=", "+", "-", "@")


class SourceUnavailable(RuntimeError):
    """The stored original is missing or no longer matches its SHA-256."""


def set_text(cell, value) -> None:
    """Write a value; text is always stored as text, never as a formula."""
    cell.value = value
    if isinstance(value, str) and value.startswith(_FORMULA):
        cell.data_type = "s"


@contextmanager
def _original_path(db: Session, report: Report):
    try:
        with get_store().local_copy(report.storage_key or "", report.source_sha256 or "", db=db) as path:
            yield Path(path)
    except (IntegrityError, FileNotFoundError) as exc:
        raise SourceUnavailable("The stored original is not available, so this file cannot be built.") from exc


def load_original(db: Session, report: Report) -> tuple[openpyxl.Workbook, bool]:
    """The customer's workbook as an editable openpyxl workbook, and whether
    its layout is the original's (xlsx/xlsm) or a faithful copy of the values
    (csv/xls, which have no layout to keep)."""
    kind = (report.file_kind or Path(report.file_name).suffix.lstrip(".")).lower()
    with _original_path(db, report) as path:
        if kind in ("xlsx", "xlsm"):
            return openpyxl.load_workbook(path), True
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        if kind == "csv":
            ws = wb.create_sheet(_sheet_title(Path(report.file_name).stem))
            text = path.read_bytes().decode("utf-8-sig", errors="replace")
            for r, row in enumerate(csv.reader(io.StringIO(text)), start=1):
                for c, v in enumerate(row, start=1):
                    if v != "":
                        set_text(ws.cell(row=r, column=c), v)
            return wb, False
        import xlrd  # .xls

        book = xlrd.open_workbook(str(path))
        for sh in book.sheets():
            ws = wb.create_sheet(_sheet_title(sh.name))
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    v = sh.cell_value(r, c)
                    if v not in ("", None):
                        set_text(ws.cell(row=r + 1, column=c + 1), v)
        return wb, False


def _sheet_title(name: str) -> str:
    return re.sub(r"[\[\]\*\?/\\:]", "_", name)[:31] or "Sheet1"


@dataclass
class Finding:
    status: str
    rule: str
    severity: str
    field_code: str | None
    cell: str | None
    column: str | None
    sentence: str
    owner: str
    claim_ref: str | None
    row: int | None
    label: str = ""
    fix: str = ""


@dataclass
class SheetFacts:
    sheet: Sheet
    claim_rows: list[int] = field(default_factory=list)  # Excel row numbers of claim rows
    excluded: list[tuple[int, int, str]] = field(default_factory=list)  # (first row, count, reason)
    columns: dict[str, str] = field(default_factory=dict)  # field_code -> source column name
    findings: dict[int, list[Finding]] = field(default_factory=lambda: defaultdict(list))  # by Excel row


def report_facts(db: Session, report: Report) -> dict[str, SheetFacts]:
    sheets = {s.id: s for s in db.query(Sheet).filter(Sheet.report_id == report.id)}
    facts = {s.sheet_name: SheetFacts(s) for s in sheets.values() if s.status == "CONFIRMED"}
    for m in db.query(Mapping).filter(Mapping.report_id == report.id, Mapping.source_column.isnot(None)):
        s = sheets.get(m.sheet_id)
        if s and s.sheet_name in facts:
            facts[s.sheet_name].columns[m.field_code] = m.source_column
    for sid, row in db.query(ClaimRow.sheet_id, ClaimRow.source_row_number).filter(ClaimRow.report_id == report.id):
        s = sheets.get(sid)
        if s and s.sheet_name in facts and row is not None:
            facts[s.sheet_name].claim_rows.append(int(row))
    for e in db.query(ExcludedRow).filter(ExcludedRow.report_id == report.id):
        if e.sheet_name in facts:
            facts[e.sheet_name].excluded.append((e.row_number, e.row_count or 1, e.reason))
    q = (db.query(ValidationResult, ClaimRow.sheet_id, ClaimRow.source_row_number, ClaimRow.claim_reference)
         .join(ClaimRow, ClaimRow.id == ValidationResult.claim_row_id)
         .filter(ValidationResult.report_id == report.id))
    for vr, sid, row, ref in q:
        s = sheets.get(sid)
        if not s or s.sheet_name not in facts:
            continue
        x = vr.extra or {}
        r = catalogue_rule(vr.rule or "")
        f = Finding(status=vr.status, rule=vr.rule or "", severity=vr.severity, field_code=x.get("field_code"),
                    cell=x.get("cell"), column=x.get("column"), sentence=x.get("sentence") or vr.message,
                    owner=x.get("owner") or r.owner, claim_ref=ref, row=int(row) if row is not None else None,
                    label=r.label if vr.status != "NOT_EVALUABLE" else "Couldn't check",
                    fix=r.fix if vr.status != "NOT_EVALUABLE" else "Map the missing column, or ask the sender for it.")
        if f.row is not None:
            facts[s.sheet_name].findings[f.row].append(f)
    return facts


def _ws_for(wb: openpyxl.Workbook, name: str):
    if name in wb.sheetnames:
        return wb[name]
    t = _sheet_title(name)
    return wb[t] if t in wb.sheetnames else None


def _col_index(sf: SheetFacts, field_code: str | None) -> int | None:
    col = sf.columns.get(field_code or "")
    headers = [str(h) for h in (sf.sheet.headers or [])]
    return headers.index(col) + 1 if col in headers else None


def _location(sheet: str, row: int | None, col: int | None) -> tuple[str, str]:
    """(display text, in-workbook link target)."""
    quoted = "'" + sheet.replace("'", "''") + "'"
    if row is None:
        return sheet, f"{quoted}!A1"
    if col is None:
        return f"{sheet}!{row}:{row}", f"{quoted}!A{row}"
    ref = f"{get_column_letter(col)}{row}"
    return f"{sheet}!{ref}", f"{quoted}!{ref}"


# ----------------------------------------------------------------- annotated workbook

def annotated_workbook(db: Session, report: Report) -> bytes:
    wb, _layout_kept = load_original(db, report)
    facts = report_facts(db, report)
    issues: list[tuple[Finding, str, str]] = []  # (finding, display location, link target)
    comments = 0
    for name, sf in facts.items():
        ws = _ws_for(wb, name)
        if ws is None:
            continue
        ncols = max(len(sf.sheet.headers or []), 1)
        header_row = (sf.sheet.header_row_index or 0) + 1
        for c in range(1, ncols + 1):
            ws.cell(row=header_row, column=c).fill = FILL["grey"]
        for first, count, _reason in sf.excluded:
            for r in range(first, first + min(count, 1000)):
                for c in range(1, ncols + 1):
                    ws.cell(row=r, column=c).fill = FILL["grey"]
        notes_col = ws.max_column + 1
        set_text(ws.cell(row=header_row, column=notes_col), "Review notes")
        ws.cell(row=header_row, column=notes_col).font = BOLD
        ws.cell(row=header_row, column=notes_col).fill = FILL["grey"]
        ws.column_dimensions[get_column_letter(notes_col)].width = 60
        for r in sf.claim_rows:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = FILL["ok"]
            found = sorted(sf.findings.get(r, []), key=lambda f: (0 if f.status == "FAIL" else 1,
                                                                   SEVERITY_ORDER.get(f.severity, 9)))
            note = ws.cell(row=r, column=notes_col)
            if not found:
                set_text(note, "OK")
                note.fill = FILL["ok"]
                continue
            set_text(note, " | ".join(f"{RESULT.get(f.status, f.status)}: {f.sentence}" for f in found)[:32000])
            note.fill = FILL[TONE_OF.get(found[0].status, "warn")]
            if found[0].status == "FAIL":
                note.font = DARK_RED
            by_cell: dict[int, list[Finding]] = defaultdict(list)
            for f in found:
                col = _col_index(sf, f.field_code)
                text, link = _location(name, r, col)
                issues.append((f, text, link))
                if col is not None:
                    by_cell[col].append(f)
            for col, fs in by_cell.items():
                cell = ws.cell(row=r, column=col)
                worst = "FAIL" if any(f.status == "FAIL" for f in fs) else fs[0].status
                cell.fill = FILL[TONE_OF.get(worst, "warn")]
                if worst == "FAIL":
                    cell.font = Font(color=RED_TEXT, bold=cell.font.bold if cell.font else False)
                if comments < MAX_COMMENTS:
                    body = "\n".join(f"{f.sentence}\nFix: {f.fix}" for f in fs)[:2000]
                    cell.comment = Comment(body, "TrueBind", width=320, height=140)
                    comments += 1
    _issues_sheet(wb, issues)
    _summary_sheet(wb, report, issues, comments >= MAX_COMMENTS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _issues_sheet(wb: openpyxl.Workbook, issues: list[tuple[Finding, str, str]]) -> None:
    ws = wb.create_sheet("Issues")
    head = ["Issue", "Result", "Severity", "Where (click to go)", "Claim reference", "What we found", "Fix",
            "Who fixes"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(row=1, column=c).font = BOLD
        ws.cell(row=1, column=c).fill = FILL["grey"]
    ws.sheet_properties.outlinePr.summaryBelow = False
    groups: dict[str, list[tuple[Finding, str, str]]] = defaultdict(list)
    for item in issues:
        groups[item[0].label].append(item)
    order = sorted(groups, key=lambda k: (0 if groups[k][0][0].status == "FAIL" else 1 if groups[k][0][0].status == "REVIEW"
                                          else 2, SEVERITY_ORDER.get(groups[k][0][0].severity, 9), k))
    r = 1
    for label in order:
        items = groups[label]
        r += 1
        set_text(ws.cell(row=r, column=1), f"{label} ({len(items)})")
        ws.cell(row=r, column=1).font = BOLD
        for f, text, link in items:
            r += 1
            values = [label, RESULT.get(f.status, f.status), f.severity.title(), text, f.claim_ref, f.sentence, f.fix,
                      "Sender" if f.owner == "sender" else "Us"]
            for c, v in enumerate(values, start=1):
                set_text(ws.cell(row=r, column=c), v)
            ws.cell(row=r, column=2).fill = FILL[TONE_OF.get(f.status, "warn")]
            where = ws.cell(row=r, column=4)
            where.hyperlink = Hyperlink(ref=where.coordinate, location=link, display=text)
            where.font = Font(color="0563C1", underline="single")
            ws.row_dimensions[r].outline_level = 1
    for col, width in zip("ABCDEFGH", (34, 14, 10, 22, 18, 70, 60, 10)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def _summary_sheet(wb: openpyxl.Workbook, report: Report, issues, comments_capped: bool) -> None:
    ws = wb.create_sheet("Review Summary", 0)
    wb.active = 0
    summary = report.summary if isinstance(report.summary, dict) else {}
    hv = summary.get("health_view") or {}
    rows = [
        ["TrueBind review", None],
        ["File", report.file_name],
        ["Verdict", hv.get("verdict_label", "")],
        [None, None],
        ["Count", "Findings"],
        ["Errors", '=COUNTIFS(Issues!$B:$B,"Error")'],
        ["Warnings", '=COUNTIFS(Issues!$B:$B,"Warning")'],
        ["Couldn't check", '=COUNTIFS(Issues!$B:$B,"Couldn\'t check")'],
        [None, None],
        ["Issue", "Errors", "Warnings", "Couldn't check"],
    ]
    for r in rows:
        ws.append([None] * len(r))
        for c, v in enumerate(r, start=1):
            if v is not None:
                ws.cell(row=ws.max_row, column=c).value = v  # formulas here are ours, never cell text
    ws["A1"].font = Font(bold=True, size=14)
    for r in (5, 10):
        for c in range(1, 5):
            ws.cell(row=r, column=c).font = BOLD
    labels = sorted({f.label for f, _, _ in issues})
    for label in labels:
        ws.append([None, None, None, None])
        r = ws.max_row
        set_text(ws.cell(row=r, column=1), label)
        esc = label.replace('"', '""')
        ws.cell(row=r, column=2).value = f'=COUNTIFS(Issues!$A:$A,"{esc}",Issues!$B:$B,"Error")'
        ws.cell(row=r, column=3).value = f'=COUNTIFS(Issues!$A:$A,"{esc}",Issues!$B:$B,"Warning")'
        ws.cell(row=r, column=4).value = f'=COUNTIFS(Issues!$A:$A,"{esc}",Issues!$B:$B,"Couldn\'t check")'
    ws.append([None])
    ws.append(["Legend"])
    ws.cell(row=ws.max_row, column=1).font = BOLD
    for key, text in (("ok", "Checked: no issue found"),
                      ("warn", "Unusual, or could not be validated (for example a date held as text)"),
                      ("err", "Error: must be fixed before the file is sent"),
                      ("grey", "Structural row (header, title, subtotal, blank), not a claim")):
        ws.append(["", text])
        ws.cell(row=ws.max_row, column=1).fill = FILL[key]
        if key == "err":
            ws.cell(row=ws.max_row, column=1).font = DARK_RED
            set_text(ws.cell(row=ws.max_row, column=1), "Error")
    ws.append([None])
    ws.append(["Hover over an amber or red cell to see what was found and how to fix it. Every row also has a "
               "Review notes column. Values in your sheets are exactly as received."])
    if comments_capped:
        ws.append([f"Hover notes were added to the first {MAX_COMMENTS:,} flagged cells; every finding is on the "
                   "Issues sheet."])
    ws.column_dimensions["A"].width = 40
    for col in "BCD":
        ws.column_dimensions[col].width = 16


# ----------------------------------------------------------------- corrected copy ("mend")

_ISO_DATE = re.compile(r"^\s*(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ T]00:00(?::00)?)?\s*$")
_DMY = re.compile(r"^\s*(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})\s*$")
_MONTH_NAME = re.compile(r"[A-Za-z]{3,}")
_HIDDEN = re.compile("[​‌‍﻿]")
_DATE_FIELDS = {"CR0119CM": "date_of_loss", "CR0136CM": "date_notified", "TB_INCEPTION": "policy_inception",
                "TB_EXPIRY": "policy_expiry"}


def _unambiguous_date(raw, parsed: dt.date | None) -> dt.date | None:
    """The engine's reading of a date held as text or a serial number, but only
    when no other reading is possible: ISO order, a month name, an Excel serial,
    or day/month where one part is above 12."""
    if parsed is None or raw is None:
        return None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return parsed if 1 <= raw <= 2958465 else None
    s = str(raw)
    if _ISO_DATE.match(s) or _MONTH_NAME.search(s):
        return parsed
    m = _DMY.match(s)
    if m and (int(m.group(1)) > 12 or int(m.group(2)) > 12):
        return parsed
    return None


@dataclass
class Change:
    sheet: str
    cell: str
    old: object
    new: object
    rule: str


def corrected_workbook(db: Session, report: Report) -> tuple[bytes, list[Change]]:
    wb, _ = load_original(db, report)
    facts = report_facts(db, report)
    rows_by_sheet: dict[str, dict[int, ClaimRow]] = defaultdict(dict)
    sheet_names = {s.sheet.id: n for n, s in facts.items()}
    for cr in db.query(ClaimRow).filter(ClaimRow.report_id == report.id):
        if cr.sheet_id in sheet_names and cr.source_row_number is not None:
            rows_by_sheet[sheet_names[cr.sheet_id]][int(cr.source_row_number)] = cr
    changes: list[Change] = []
    valid_status = {v.lower() for v in FIELDS_BY_CODE["CR0105CM"].enum_values}

    for name, sf in facts.items():
        ws = _ws_for(wb, name)
        if ws is None:
            continue
        col = {code: _col_index(sf, code) for code in sf.columns}
        pol_col = col.get("CR0029M")
        pad_to = _policy_pad_width(ws, pol_col, sf.claim_rows) if pol_col else None

        def change(r: int, c: int, new, why: str):
            cell = ws.cell(row=r, column=c)
            if cell.value == new:
                return
            changes.append(Change(name, cell.coordinate, cell.value, new, why))
            if isinstance(new, str):
                set_text(cell, new)
            else:
                cell.value = new
            if isinstance(new, dt.date):
                cell.number_format = "yyyy-mm-dd"

        for r in sf.claim_rows:
            cr = rows_by_sheet[name].get(r)
            rules = {(f.rule, f.field_code) for f in sf.findings.get(r, [])}
            # trimmed spaces / hidden characters on every mapped text cell
            for code, c in col.items():
                if c is None:
                    continue
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str):
                    t = _HIDDEN.sub("", v).replace(" ", " ").strip()
                    t = re.sub(r" {2,}", " ", t)
                    if t != v and t:
                        change(r, c, t, "trimmed spaces / removed hidden characters")
            if cr is None:
                continue
            for code, attr in _DATE_FIELDS.items():
                c = col.get(code)
                if c and ("date_stored_as_text", code) in rules:
                    d = _unambiguous_date(ws.cell(row=r, column=c).value, getattr(cr, attr, None))
                    if d is not None:
                        change(r, c, d, "date held as text or a serial number -> real date")
            c = col.get("CR0110CM")
            if c and ("currency_normalised", "CR0110CM") in rules and cr.currency:
                change(r, c, cr.currency, "currency written non-standardly -> ISO code")
            c = col.get("CR0105CM")
            v = ws.cell(row=r, column=c).value if c else None
            if c and isinstance(v, str) and cr.claim_status and cr.claim_status.lower() in valid_status:
                key = " ".join(v.strip().split()).lower()
                if key in CLAIM_STATUS_SYNONYMS or (key == cr.claim_status.lower() and v.strip() != v.strip().title()):
                    change(r, c, cr.claim_status.title(), "status synonym -> agreed status")
            for code, c in col.items():
                if c and ("amount_stored_as_text", code) in rules:
                    attr = _AMOUNT_FIELDS.get(code)
                    val = getattr(cr, attr, None) if attr else None
                    if val is not None:
                        change(r, c, float(val), "amount held as text -> number")
            if pad_to and pol_col:
                v = ws.cell(row=r, column=pol_col).value
                s = str(int(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() \
                    else (v.strip() if isinstance(v, str) else None)
                if s and s.isdigit() and len(s) < pad_to:
                    change(r, pol_col, s.zfill(pad_to), f"policy number zero-padded to {pad_to} digits")

    log = wb.create_sheet("Change Log", 0)
    log.append(["Sheet", "Cell", "Old value", "New value", "Rule"])
    for c in range(1, 6):
        log.cell(row=1, column=c).font = BOLD
        log.cell(row=1, column=c).fill = FILL["grey"]
    for ch in changes:
        log.append([None] * 5)
        r = log.max_row
        for c, v in enumerate((ch.sheet, ch.cell, _show(ch.old), _show(ch.new), ch.rule), start=1):
            set_text(log.cell(row=r, column=c), v)
    if not changes:
        log.append(["No safe fixes were needed."])
    log.append([None])
    log.append(["Only safe, reversible fixes are made. Duplicates, arithmetic, reserves and anything ambiguous are "
                "never changed: they are in the query letter for the sender."])
    for col_letter, width in zip("ABCDE", (24, 10, 34, 34, 48)):
        log.column_dimensions[col_letter].width = width
    log.cell(row=1, column=1).alignment = Alignment(vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), changes


_AMOUNT_FIELDS = {"TB_PAID_TD": "paid_amount", "CR0130CM": "reserve_amount", "CR0155CM": "incurred_amount",
                  "CR0134CM": "incurred_indemnity", "CR0126CM": "paid_this_month", "CR0128CM": "previously_paid",
                  "CR0127CM": "fees_paid_this_month", "CR0129CM": "fees_previously_paid", "CR0131CM": "fees_reserve",
                  "TB_FEES_PAID_TD": "fees_paid_to_date", "TB_POLICY_LIMIT": "policy_limit"}


def _show(v) -> str:
    if v is None:
        return "(blank)"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10]
    return str(v)


def _policy_pad_width(ws, col: int, rows: list[int]) -> int | None:
    """Zero-pad only when nearly every policy number is all digits of one
    length (>= 80%, at least 10 rows): then a shorter one lost its zeros."""
    lengths: dict[int, int] = defaultdict(int)
    n = 0
    for r in rows:
        v = ws.cell(row=r, column=col).value
        s = str(int(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() \
            else (v.strip() if isinstance(v, str) else "")
        if s.isdigit():
            lengths[len(s)] += 1
            n += 1
    if n < 10:
        return None
    width, count = max(lengths.items(), key=lambda kv: kv[1])
    return width if count / n >= 0.8 else None


# ----------------------------------------------------------------- query letter

def query_letter(db: Session, report: Report, facts: dict[str, SheetFacts] | None = None) -> dict:
    facts = facts if facts is not None else report_facts(db, report)
    groups: dict[str, dict] = {}
    for name, sf in facts.items():
        for r, fs in sf.findings.items():
            for f in fs:
                if f.owner != "sender" or f.status not in ("FAIL", "REVIEW"):
                    continue
                g = groups.setdefault(f.rule, {"rule": f.rule, "label": f.label, "fix": f.fix,
                                               "outcome": f.status, "severity": f.severity, "items": []})
                col = _col_index(sf, f.field_code)
                where, _ = _location(name, r, col)
                g["items"].append({"where": where, "claim_ref": f.claim_ref, "sentence": f.sentence})
    ordered = sorted(groups.values(), key=lambda g: (0 if g["outcome"] == "FAIL" else 1,
                                                     SEVERITY_ORDER.get(g["severity"], 9), -len(g["items"])))
    for g in ordered:
        g["items"].sort(key=lambda i: i["where"])
    total = sum(len(g["items"]) for g in ordered)
    who = report.sender or "team"
    subject = f"Queries on {report.file_name}"
    lines = [f"Dear {who},", "",
             f"Thank you for {report.file_name}. Before we can accept it, please check the points below "
             f"({total} item{'s' if total != 1 else ''}). Each line gives the sheet, the cell and what we found.", ""]
    if not ordered:
        lines = [f"Dear {who},", "", f"Thank you for {report.file_name}. We have no queries on this file.", ""]
    for i, g in enumerate(ordered, start=1):
        lines.append(f"{i}. {g['label']} ({len(g['items'])})")
        lines.append(f"   What we need: {g['fix']}")
        for it in g["items"][:200]:
            ref = f" (claim {it['claim_ref']})" if it["claim_ref"] else ""
            lines.append(f"   - {it['where']}{ref}: {it['sentence']}")
        if len(g["items"]) > 200:
            lines.append(f"   - and {len(g['items']) - 200} more; the full list is in the attached annotated workbook.")
        lines.append("")
    lines += ["Please send a corrected file, or reply with an explanation for any item you believe is right.", "",
              "Kind regards,"]
    return {"subject": subject, "body": "\n".join(lines), "groups": ordered, "items": total}


# ----------------------------------------------------------------- month-on-month

CLOSED = {"closed", "ntu"}


def _f(v) -> float | None:
    return None if v is None else float(v)


def month_on_month(db: Session, current: Report, previous: Report, reserve_jump_pct: float = 50.0,
                   reserve_jump_min: float = 0.0) -> dict:
    """Claims in the previous report that vanished without closing, paid to
    date that went down, and reserves that rose by more than the threshold
    (same claim reference, same currency only)."""
    def latest(report_id: str) -> dict[str, ClaimRow]:
        out: dict[str, ClaimRow] = {}
        for cr in (db.query(ClaimRow).filter(ClaimRow.report_id == report_id)
                   .order_by(ClaimRow.sheet_id, ClaimRow.row_index)):
            if cr.claim_reference:
                out[str(cr.claim_reference).strip().upper()] = cr
        return out

    sheet_names = {s.id: s.sheet_name for s in db.query(Sheet).filter(Sheet.report_id.in_([current.id, previous.id]))}
    prev, cur = latest(previous.id), latest(current.id)

    def where(cr: ClaimRow) -> str:
        return f"{sheet_names.get(cr.sheet_id, '?')} row {cr.source_row_number}"

    vanished, paid_down, reserve_jump = [], [], []
    for ref, p in prev.items():
        c = cur.get(ref)
        if c is None:
            if (p.claim_status or "").lower() not in CLOSED:
                vanished.append({"claim_ref": p.claim_reference, "previous_status": p.claim_status,
                                 "previous": where(p), "sentence": f"Claim {p.claim_reference} was "
                                 f"{p.claim_status or 'open'} last time and is missing from this file without being closed."})
            continue
        if (p.currency or "") != (c.currency or ""):
            continue
        p_paid, c_paid = _f(p.paid_amount), _f(c.paid_amount)
        p_res, c_res = _f(p.reserve_amount), _f(c.reserve_amount)
        if p_paid is not None and c_paid is not None and c_paid < p_paid - 0.005:
            paid_down.append({"claim_ref": c.claim_reference, "where": where(c), "previous": p_paid,
                              "current": c_paid, "currency": c.currency,
                              "sentence": f"Paid to date fell from {p_paid:,.2f} to {c_paid:,.2f}; "
                                          "paid to date should never go down."})
        if p_res is not None and c_res is not None:
            rise = c_res - p_res
            pct = (rise / p_res * 100) if p_res > 0 else (100.0 if rise > 0 else 0.0)
            if rise > 0 and pct > reserve_jump_pct and rise >= reserve_jump_min:
                reserve_jump.append({"claim_ref": c.claim_reference, "where": where(c), "previous": p_res,
                                     "current": c_res, "currency": c.currency, "pct": round(pct, 1),
                                     "sentence": f"Reserve rose {pct:,.0f}% from {p_res:,.2f} to "
                                                 f"{c_res:,.2f}."})
    return {"current_report_id": current.id, "previous_report_id": previous.id,
            "previous_file_name": previous.file_name, "threshold_pct": reserve_jump_pct,
            "threshold_min": reserve_jump_min, "claims_compared": len(set(prev) & set(cur)),
            "vanished": vanished, "paid_down": paid_down, "reserve_jump": reserve_jump,
            "counts": {"vanished": len(vanished), "paid_down": len(paid_down), "reserve_jump": len(reserve_jump)}}
