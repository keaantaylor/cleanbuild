"""The workbook as the review surface: the source sheet's cells served as
tiles (a block of rows x a block of columns), each cell with its value, its
formula when it has one, its reconciliation state and the issues and
corrections on it.

Values and formulas are read from the stored original once and cached (gzip
JSON, derived from the file's SHA-256), so scrolling a 100,000-row or
500-column sheet never re-reads the workbook. States come from the stored
findings, located by the cell each finding records.

Cell state, always shown with its words, never colour alone:
  verified                 a claim cell no check flagged
  requires_reconciliation  a check failed on this cell
  undetermined             flagged for review, or the check could not run
"""

from __future__ import annotations

import datetime as dt
import gzip
import json
import re
import threading
from collections import defaultdict
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from ..models.corrections import Correction
from ..models.reports import ClaimRow, ExcludedRow, Report, Sheet, ValidationResult
from . import issues as issue_svc
from .deliverables import _original_path, _ws_for, load_original
from .storage import derived_key, get_store

MAX_COLS = 1000
MAX_CELL = 200
CELL = re.compile(r"^([A-Z]{1,3})(\d+)$")
STATE = {"FAIL": "requires_reconciliation", "REVIEW": "undetermined", "NOT_EVALUABLE": "undetermined"}
TONE = {"requires_reconciliation": "err", "undetermined": "warn", "verified": "ok"}


def _show(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:19].replace("T", " ").removesuffix(" 00:00:00")
    s = str(v)
    return s if len(s) <= MAX_CELL else s[: MAX_CELL - 1] + "…"


def col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def col_number(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


_build_locks: dict[str, threading.Lock] = defaultdict(threading.Lock)


def _sheet_cache(db: Session, report: Report, sheet_names: list[str]) -> dict[str, dict]:
    """{sheet: {"v": rows of display values, "f": {"row,col": formula}}}.
    Built once per file: the grid asks for several tiles at once, and without
    the lock each request would re-read the workbook in parallel."""
    key = derived_key(report.tenant_id, report.source_sha256 or "0" * 64, "grid-v2")
    store = get_store()

    def cached() -> dict | None:
        blob = store.get_derived(key, db=db)
        if blob is not None:
            data = json.loads(gzip.decompress(blob))
            if all(n in data for n in sheet_names):
                return data
        return None

    hit = cached()
    if hit is not None:
        return hit
    with _build_locks[key]:
        hit = cached()  # built by another request while this one waited
        if hit is not None:
            return hit
        return _build(db, report, sheet_names, store, key)


def _build(db: Session, report: Report, sheet_names: list[str], store, key: str) -> dict[str, dict]:
    wb, layout_kept = load_original(db, report)
    computed = None
    if layout_kept:  # xlsx/xlsm: the cached results of formulas, as Excel last calculated them
        with _original_path(db, report) as path:
            computed = openpyxl.load_workbook(Path(path), data_only=True)
    out: dict[str, dict] = {}
    for name in sheet_names:
        ws = _ws_for(wb, name)
        if ws is None:
            continue
        cws = computed[ws.title] if computed is not None and ws.title in computed.sheetnames else None
        # Read only up to the sheet's own last column: asking openpyxl for MAX_COLS on every
        # row materialises empty cells (5,000 rows x 1,000 columns took minutes).
        width = max(1, min(MAX_COLS, ws.max_column or 1))
        vals = [list(r) for r in cws.iter_rows(max_col=width, values_only=True)] if cws is not None else None
        rows, formulas = [], {}
        for r, row in enumerate(ws.iter_rows(max_col=width), start=1):
            line = []
            for c, cell in enumerate(row, start=1):
                v = cell.value
                if cell.data_type == "f" and isinstance(v, str):
                    formulas[f"{r},{c}"] = v[:MAX_CELL]
                    got = vals[r - 1][c - 1] if vals is not None and r - 1 < len(vals) and c - 1 < len(vals[r - 1]) else None
                    v = got if got is not None else v
                line.append(_show(v))
            while line and line[-1] is None:
                line.pop()
            rows.append(line)
        while rows and not any(rows[-1]):
            rows.pop()
        out[name] = {"v": rows, "f": formulas}
    if computed is not None:
        computed.close()
    store.put_derived(key, gzip.compress(json.dumps(out).encode()), db=db)
    return out


def _cached_values(db: Session, report: Report, sheet_names: list[str]) -> dict[str, list[list[str | None]]]:
    """{sheet: rows of display values}: used to read a cell's original value."""
    return {k: v["v"] for k, v in _sheet_cache(db, report, sheet_names).items()}


def _confirmed(db: Session, report: Report, sheet: Sheet) -> list[str]:
    names = [s.sheet_name for s in db.query(Sheet).filter(Sheet.report_id == report.id, Sheet.status == "CONFIRMED")]
    return names if sheet.sheet_name in names else [*names, sheet.sheet_name]


def _issue_cells(db: Session, report: Report, sheet: Sheet) -> tuple[dict, dict]:
    """(row, col) -> issues on that cell; row -> issues on the whole row."""
    by_cell: dict[tuple[int, int], list[dict]] = defaultdict(list)
    by_row: dict[int, list[dict]] = defaultdict(list)
    for vr in (db.query(ValidationResult).join(ClaimRow, ClaimRow.id == ValidationResult.claim_row_id)
               .filter(ValidationResult.report_id == report.id, ClaimRow.sheet_id == sheet.id)):
        x = vr.extra or {}
        cell = str(x.get("cell") or "")
        status = x.get("issue_status") or issue_svc.initial_status(vr.status, vr.rule or "")
        item = {"issue_id": vr.id, "rule": vr.rule, "rule_version": x.get("rule_version"), "outcome": vr.status,
                "status": status, "state": STATE.get(vr.status, "undetermined"), "label": x.get("sentence") or vr.message,
                "expected": x.get("expected"), "actual": x.get("actual"), "difference": x.get("difference"),
                "known_exception": x.get("known_exception")}
        m = CELL.match(cell)
        if m:
            by_cell[(int(m.group(2)), col_number(m.group(1)))].append(item)
        elif cell.startswith("row "):
            by_row[int(cell[4:])].append(item)
    return by_cell, by_row


def _corrections(db: Session, report: Report, sheet: Sheet) -> dict[tuple[int, int], dict]:
    out = {}
    for c in (db.query(Correction).filter(Correction.report_id == report.id, Correction.sheet_name == sheet.sheet_name,
                                          Correction.status != "REJECTED").order_by(Correction.created_at)):
        m = CELL.match(c.cell or "")
        if m:
            out[(int(m.group(2)), col_number(m.group(1)))] = {
                "id": c.id, "status": c.status, "after": c.after_value, "policy": c.policy,
                "verified": (c.result or {}).get("passed")}
    return out


def _cell_state(kind: str, issues: list[dict]) -> str | None:
    open_issues = [i for i in issues if i["status"] not in issue_svc.CLOSED]
    if open_issues:
        return ("requires_reconciliation" if any(i["state"] == "requires_reconciliation" for i in open_issues)
                else "undetermined")
    if issues or kind == "claim":
        return "verified"
    return None


def grid_page(db: Session, report: Report, sheet: Sheet, offset: int, limit: int, col_offset: int = 0,
              col_limit: int = MAX_COLS, rows: list[int] | None = None) -> dict:
    """A tile: rows offset+1..offset+limit (or the given row numbers) and
    columns col_offset+1..col_offset+col_limit."""
    cache = _sheet_cache(db, report, _confirmed(db, report, sheet)).get(sheet.sheet_name, {"v": [], "f": {}})
    values, formulas = cache["v"], cache["f"]
    ncols = min(max((len(r) for r in values), default=0), MAX_COLS)
    header_row = (sheet.header_row_index or 0) + 1
    wanted = sorted({r for r in rows if 1 <= r <= len(values)}) if rows is not None else \
        list(range(offset + 1, min(offset + limit, len(values)) + 1))
    c0, c1 = col_offset + 1, min(col_offset + col_limit, ncols)
    lo, hi = (wanted[0], wanted[-1]) if wanted else (0, -1)
    claim_rows = {int(r) for (r,) in db.query(ClaimRow.source_row_number).filter(
        ClaimRow.report_id == report.id, ClaimRow.sheet_id == sheet.id, ClaimRow.source_row_number.between(lo, hi))}
    structural = set()
    for e in db.query(ExcludedRow).filter(ExcludedRow.report_id == report.id, ExcludedRow.sheet_name == sheet.sheet_name):
        for r in range(e.row_number, e.row_number + min(e.row_count or 1, 100000)):
            if lo <= r <= hi:
                structural.add(r)
    by_cell, by_row = _issue_cells(db, report, sheet) if sheet.status == "CONFIRMED" else ({}, {})
    corr = _corrections(db, report, sheet)
    out_rows = []
    for r in wanted:
        cells = values[r - 1] if r - 1 < len(values) else []
        kind = ("header" if r == header_row else "structural" if r in structural or r < header_row
                else "claim" if r in claim_rows else "other")
        out_cells = []
        for c in range(c0, c1 + 1):
            cell_issues = by_cell.get((r, c), [])
            state = _cell_state(kind, cell_issues) if kind in ("claim", "structural") and (cell_issues or kind == "claim") else None
            out_cells.append({"v": cells[c - 1] if c - 1 < len(cells) else None, "f": formulas.get(f"{r},{c}"),
                              "state": state, "tone": "grey" if kind == "header" else TONE.get(state or "", None),
                              "issues": cell_issues, "correction": corr.get((r, c)),
                              # kept for the older preview
                              "notes": [{"result": i["state"], "status": i["outcome"], "label": i["label"],
                                         "text": i["label"], "fix": i["rule"]} for i in cell_issues]})
        out_rows.append({"row": r, "kind": kind, "cells": out_cells, "row_issues": by_row.get(r, []),
                         "row_notes": [{"result": i["state"], "status": i["outcome"], "label": i["label"],
                                        "text": i["label"], "fix": i["rule"]} for i in by_row.get(r, [])]})
    return {"sheet_id": sheet.id, "sheet_name": sheet.sheet_name, "total_rows": len(values), "columns": ncols,
            "header_row": header_row, "offset": offset, "col_offset": col_offset, "rows": out_rows,
            "headers": [values[header_row - 1][c - 1] if header_row - 1 < len(values) and c - 1 < len(values[header_row - 1]) else None
                        for c in range(1, ncols + 1)] if ncols <= MAX_COLS else []}


def search(db: Session, report: Report, sheet: Sheet, q: str, limit: int = 200) -> list[dict]:
    """Cells whose value or formula contains q (case-insensitive)."""
    cache = _sheet_cache(db, report, _confirmed(db, report, sheet)).get(sheet.sheet_name, {"v": [], "f": {}})
    needle = q.strip().lower()
    hits = []
    if not needle:
        return hits
    for r, row in enumerate(cache["v"], start=1):
        for c, v in enumerate(row, start=1):
            f = cache["f"].get(f"{r},{c}")
            if (v is not None and needle in v.lower()) or (f and needle in f.lower()):
                hits.append({"row": r, "col": c, "cell": f"{col_letter(c)}{r}", "v": v})
                if len(hits) >= limit:
                    return hits
    return hits


def issue_cells(db: Session, report: Report, sheet: Sheet) -> list[dict]:
    """Every flagged cell on the sheet, in reading order, for 'issues only'
    filtering and next / previous issue navigation."""
    by_cell, by_row = _issue_cells(db, report, sheet)
    out = [{"row": r, "col": c, "cell": f"{col_letter(c)}{r}", "state": _cell_state("claim", items),
            "open": sum(1 for i in items if i["status"] not in issue_svc.CLOSED),
            "issue_ids": [i["issue_id"] for i in items]} for (r, c), items in by_cell.items()]
    out += [{"row": r, "col": None, "cell": f"row {r}", "state": _cell_state("claim", items),
             "open": sum(1 for i in items if i["status"] not in issue_svc.CLOSED),
             "issue_ids": [i["issue_id"] for i in items]} for r, items in by_row.items()]
    return sorted(out, key=lambda x: (x["row"], x["col"] or 0))
