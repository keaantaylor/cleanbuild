"""Synthetic claims bordereau with planted problems and an answer key.

Every planted problem is recorded in the answer key with its sheet, Excel row,
the rule TrueBind should report and the cell it concerns, so a run can be
scored for recall (planted problems found) and false positives (findings on
rows where nothing was planted). Deterministic for a given seed.

    python tests/fixtures/generate_bordereau.py --rows 2000 --out out.xlsx [--seed 7]
writes out.xlsx and out.answers.json.

The workbook mimics a coverholder's monthly file: a title row, a header row,
claim rows, a subtotal row, a Summary tab and a Notes tab (both non-claims).
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import random
from pathlib import Path

import openpyxl

HEADER = ["Claim Ref", "Pol No.", "Insured Name", "Inception", "Expiry", "Policy Limit", "Date of Loss", "Date Rptd",
          "Claim Status", "Currency", "Paid to Date", "Outstanding Reserve", "Total Incurred"]
COL = {h: i + 1 for i, h in enumerate(HEADER)}
LETTER = {h: openpyxl.utils.get_column_letter(i) for h, i in COL.items()}

FIRST = ["Harbour", "Kestrel", "Norland", "Ashby", "Brennan", "Calder", "Dunmore", "Ellery", "Fenwick", "Galway",
         "Hollis", "Inglewood", "Juniper", "Kilbride", "Larkspur", "Merrow", "Northcote", "Oakham", "Pembroke",
         "Quarry", "Rathmore", "Selby", "Thornbury", "Upton", "Vale", "Westbury", "Yarrow", "Ardmore", "Ballina"]
SECOND = ["Freight", "Logistics", "Bakery", "Dental", "Motors", "Marine", "Builders", "Textiles", "Farms", "Hotels",
          "Pharmacy", "Engineering", "Print", "Joinery", "Haulage", "Foods", "Garages", "Electrical", "Florists"]
SUFFIX = ["Ltd", "Limited", "plc", "& Sons", "Group", "LLP", "Holdings"]

PLANTS = [  # (kind, rule the engine should report, header of the cell concerned, outcome)
    ("arithmetic", "arithmetic_mismatch", "Total Incurred", "FAIL"),
    ("missing_name", "missing_mandatory_field", "Insured Name", "FAIL"),
    ("bad_currency", "invalid_currency", "Currency", "FAIL"),
    ("currency_variant", "currency_normalised", "Currency", "REVIEW"),
    ("notified_before_loss", "date_order", "Date Rptd", "FAIL"),
    ("loss_outside_period", "loss_outside_policy_period", "Date of Loss", "FAIL"),
    ("over_limit", "incurred_over_limit", "Total Incurred", "FAIL"),
    ("negative_reserve", "negative_reserve", "Outstanding Reserve", "FAIL"),
    ("closed_with_reserve", "closed_with_reserve", "Outstanding Reserve", "REVIEW"),
    ("bad_status", "invalid_status", "Claim Status", "FAIL"),
    ("date_text", "date_stored_as_text", "Date of Loss", "REVIEW"),
    ("exact_duplicate", "exact_duplicate", "Claim Ref", "FAIL"),
    ("probable_duplicate", "probable_duplicate", "Claim Ref", "REVIEW"),
]


def _name(rng: random.Random) -> str:
    return f"{rng.choice(FIRST)} {rng.choice(SECOND)} {rng.choice(SUFFIX)}"


def generate(rows: int, out: Path, seed: int = 7, rate: float = 0.03) -> dict:
    rng = random.Random(seed)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Claims"
    ws.append(["Monthly claims bordereau - synthetic sample (no real data)"])
    ws.append(HEADER)
    answers: list[dict] = []
    clean_rows: list[int] = []
    planted_rows: set[int] = set()
    today = dt.date(2026, 9, 30)
    data: list[list] = []

    for i in range(rows):
        inception = dt.date(2025, rng.randint(1, 12), rng.randint(1, 28))
        expiry = inception + dt.timedelta(days=364)
        loss = inception + dt.timedelta(days=rng.randint(5, 330))
        loss = min(loss, today - dt.timedelta(days=10))
        reported = loss + dt.timedelta(days=rng.randint(1, 40))
        reported = min(reported, today)
        status = rng.choices(["Open", "Closed", "Reopened"], [70, 25, 5])[0]
        limit = rng.choice([100000, 250000, 500000, 1000000])
        paid = round(rng.uniform(0, min(60000, 0.45 * limit)), 2)  # clean rows stay within the limit
        reserve = 0.0 if status == "Closed" else round(rng.uniform(500, min(80000, 0.45 * limit)), 2)
        data.append([f"CLM-{100000 + i * 7}", f"{4000000 + rng.randint(0, 899999)}", _name(rng), inception, expiry,
                     limit, loss, reported, status, rng.choice(["GBP", "GBP", "GBP", "EUR", "USD"]), paid, reserve,
                     round(paid + reserve, 2)])

    def plant(idx: int, kind: str) -> None:
        r = data[idx]
        if kind == "arithmetic":
            r[COL["Total Incurred"] - 1] = round(r[COL["Total Incurred"] - 1] + rng.choice([-1, 1]) * rng.uniform(250, 5000), 2)
        elif kind == "missing_name":
            r[COL["Insured Name"] - 1] = None
        elif kind == "bad_currency":
            r[COL["Currency"] - 1] = rng.choice(["XXX", "GPB", "EU"])
        elif kind == "currency_variant":
            r[COL["Currency"] - 1] = rng.choice(["Euro", "eur", "€", "Sterling"])
        elif kind == "notified_before_loss":
            loss = r[COL["Date of Loss"] - 1]
            r[COL["Date Rptd"] - 1] = loss - dt.timedelta(days=rng.randint(3, 30))
        elif kind == "loss_outside_period":
            r[COL["Date of Loss"] - 1] = r[COL["Inception"] - 1] - dt.timedelta(days=rng.randint(5, 40))
            r[COL["Date Rptd"] - 1] = r[COL["Date of Loss"] - 1] + dt.timedelta(days=5)
        elif kind == "over_limit":
            r[COL["Policy Limit"] - 1] = 100000
            r[COL["Paid to Date"] - 1] = 90000.0
            r[COL["Outstanding Reserve"] - 1] = 60000.0
            r[COL["Total Incurred"] - 1] = 150000.0
        elif kind == "negative_reserve":
            r[COL["Outstanding Reserve"] - 1] = -round(rng.uniform(100, 5000), 2)
            r[COL["Total Incurred"] - 1] = round(r[COL["Paid to Date"] - 1] + r[COL["Outstanding Reserve"] - 1], 2)
            if r[COL["Total Incurred"] - 1] < r[COL["Paid to Date"] - 1]:
                pass  # paid > incurred is a consequence; scored as a related finding, not a false positive
        elif kind == "closed_with_reserve":
            r[COL["Claim Status"] - 1] = "Closed"
            if r[COL["Outstanding Reserve"] - 1] <= 0:
                r[COL["Outstanding Reserve"] - 1] = 1500.0
                r[COL["Total Incurred"] - 1] = round(r[COL["Paid to Date"] - 1] + 1500.0, 2)
        elif kind == "bad_status":
            r[COL["Claim Status"] - 1] = rng.choice(["Pending?", "TBC", "In Litigation Hold"])
        elif kind == "date_text":
            d = r[COL["Date of Loss"] - 1]
            r[COL["Date of Loss"] - 1] = f"{d.day:02d} {d.strftime('%b')} {d.year}"  # text, but unambiguous

    n_plants = max(len(PLANTS), int(rows * rate))
    candidates = list(range(rows))
    rng.shuffle(candidates)
    chosen = candidates[:n_plants]
    dup_sources: list[tuple[int, str]] = []
    for k, idx in enumerate(chosen):
        kind, rule, header, outcome = PLANTS[k % len(PLANTS)]
        if kind in ("exact_duplicate", "probable_duplicate"):
            dup_sources.append((idx, kind))
            continue
        plant(idx, kind)
        planted_rows.add(idx)
        answers.append({"kind": kind, "rule": rule, "outcome": outcome, "data_index": idx, "header": header})

    out_rows: list[list] = list(data)
    for idx, kind in dup_sources:
        copy = list(data[idx])
        if kind == "probable_duplicate":
            # Re-keyed under a new reference: same insured, policy, loss and amounts.
            copy[0] = f"CLM-{900000 + idx}"
        out_rows.append(copy)
        planted_rows.add(len(out_rows) - 1)
        planted_rows.add(idx)
        answers.append({"kind": kind, "rule": "exact_duplicate" if kind == "exact_duplicate" else "probable_duplicate",
                        "outcome": "FAIL" if kind == "exact_duplicate" else "REVIEW", "data_index": len(out_rows) - 1,
                        "pair_index": idx, "header": "Claim Ref"})

    for r in out_rows:
        ws.append(r)
    first = 3
    total_row = first + len(out_rows)
    ws.append(["Total", None, None, None, None, None, None, None, None, None,
               round(sum(r[10] or 0 for r in out_rows), 2), round(sum(r[11] or 0 for r in out_rows), 2),
               round(sum(r[12] or 0 for r in out_rows), 2)])
    for row in ws.iter_rows(min_row=first, max_row=total_row - 1):
        for c in row:
            if isinstance(c.value, dt.date):
                c.number_format = "dd/mm/yyyy"
    for a in answers:
        a["sheet"] = "Claims"
        a["row"] = first + a["data_index"]
        a["cell"] = f"{LETTER[a['header']]}{a['row']}"
        if "pair_index" in a:
            a["pair_row"] = first + a["pair_index"]

    s = wb.create_sheet("Summary")
    s.append(["Status", "Claims", "Total incurred"])
    for st in ("Open", "Closed", "Reopened"):
        s.append([st, sum(1 for r in out_rows if r[8] == st), round(sum(r[12] for r in out_rows if r[8] == st), 2)])
    n = wb.create_sheet("Notes")
    n.append(["Prepared by the claims team. Figures in policy currency."])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    key = {"file": out.name, "seed": seed, "claim_rows": len(out_rows), "first_row": first,
           "planted": answers, "planted_rows": sorted(first + i for i in planted_rows),
           "non_claims_sheets": ["Summary", "Notes"], "subtotal_row": total_row}
    out.with_suffix(".answers.json").write_text(json.dumps(key, indent=1, default=str), encoding="utf-8")
    return key


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--rows", type=int, default=500)
    ap.add_argument("--out", required=True)
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--rate", type=float, default=0.03)
    a = ap.parse_args()
    k = generate(a.rows, Path(a.out), a.seed, a.rate)
    print(f"{a.out}: {k['claim_rows']} claim rows, {len(k['planted'])} planted problems")
