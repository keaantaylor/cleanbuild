"""The TrueBind Health Report (.xlsx) and the extended annotated workbook,
built from the 320-row, 10-sheet boundary fixture and checked against its
answer key, then re-opened with openpyxl to prove the files are sound."""

from __future__ import annotations

import io
import json
import time

import openpyxl
import pytest
from conftest import BORDEREAUX_ROOT, Api

from app.exports.health_report import SHEETS

BOUNDARY = BORDEREAUX_ROOT / "data" / "synthetic" / "test_boundary_cases.xlsx"
ANSWERS = json.loads((BORDEREAUX_ROOT / "data" / "synthetic" / "test_boundary_cases_answer_key.json").read_text())
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def boundary_run():
    api = Api(email="health@a.example", org="Health Org")
    rid, report = api.full_run(BOUNDARY.name, BOUNDARY.read_bytes())
    r = api.get(f"/api/v1/reports/{rid}/export/health-report.xlsx")
    assert r.status_code == 200, r.text
    a = api.get(f"/api/v1/reports/{rid}/export/annotated.xlsx")
    assert a.status_code == 200, a.text
    return api, rid, r, a


def _table_rows(ws, header_row: int, key_col: int = 1):
    """Data rows under a header until the first row whose key cell is empty."""
    out = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[key_col - 1] in (None, ""):
            break
        out.append(row)
    return out


def _header_row(ws, first: str, start: int = 1) -> int:
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, 1).value == first:
            return r
    raise AssertionError(f"no header {first!r} on {ws.title}")


def test_health_report_download_headers(boundary_run):
    _, _, r, _ = boundary_run
    assert r.headers["content-type"].startswith(XLSX)
    cd = r.headers["content-disposition"]
    assert cd.startswith('attachment; filename="TrueBind_Health_Report_test_boundary_cases_')
    assert cd.endswith('.xlsx"')


def test_health_report_sheets_and_answer_key(boundary_run):
    _, _, r, _ = boundary_run
    wb = openpyxl.load_workbook(io.BytesIO(r.content))  # raises on a corrupt file
    assert wb.sheetnames == list(SHEETS)
    assert wb.active.title == "Summary"
    assert wb.properties.creator == "TrueBind"

    # Coverage banner: 10 of 10 sheets, and every row read.
    summary = wb["Summary"]
    texts = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str)]
    assert f"{ANSWERS['sheets_processed']} of {ANSWERS['sheets_processed']}" in texts
    assert any(t.startswith("COVERAGE:") for t in texts)

    # Exceptions: one table row per finding.
    exc = wb["Exceptions"]
    assert exc.tables["Exceptions"].ref.startswith("A1:")
    rows = _table_rows(exc, 1)
    head = [c.value for c in exc[1]]
    check = head.index("Check")
    by_check = {}
    for row in rows:
        by_check[row[check]] = by_check.get(row[check], 0) + 1
    assert by_check["Total incurred does not reconcile"] == len(ANSWERS["injected_errors"]["arithmetic_errors"]["rows"])

    # Reconciliation: 38 mismatches, each with a live difference formula.
    rec = wb["Reconciliation"]
    hr = _header_row(rec, "Ref")
    rec_rows = _table_rows(rec, hr)
    mismatches = [x for x in rec_rows if x[4] == "Total incurred does not reconcile"]
    assert len(mismatches) == 38
    assert all(str(x[13]).startswith("=IF(") for x in rec_rows)
    totals = [c.value for row in rec.iter_rows(min_row=hr + len(rec_rows) + 1) for c in row]
    assert any(isinstance(v, str) and v.startswith("=SUMIFS(") for v in totals)

    # Duplicates: the answer key plants 16 name-and-date-only pairs, which the
    # engine deliberately no longer reports without corroboration (see
    # bordereaux/tests/test_boundary_fixture.py). The sheet must agree with the
    # engine: none. Real duplicates are covered by test_duplicates_sheet_groups.
    assert len(ANSWERS["injected_errors"]["duplicates"]["pairs"]) == 16
    assert wb["Duplicates"]["A2"].value.startswith("No duplicates were found")

    # Missing data: 68 distinct rows (48 claim refs + 26 insured names, 6 overlap).
    md = wb["Missing Data"]
    assert md["A2"].value.startswith("68 distinct rows")
    hr = _header_row(md, "Ref", start=4)
    listed = _table_rows(md, hr)
    assert len({(x[1], x[2]) for x in listed}) == 68
    assert len(listed) == 48 + 26

    # Coverage lists every source column of the 10 sheets.
    cov = wb["Coverage"]
    hr = _header_row(cov, "Sheet", start=5)
    hr = _header_row(cov, "Sheet", start=hr + 1)  # the Columns section's header
    cols = _table_rows(cov, hr)
    assert {x[0] for x in cols} == set(ANSWERS["sheet_row_counts"])
    assert {x[4] for x in cols} <= {"Mapped by alias", "Mapped by AI", "Mapped from memory", "Mapped by hand",
                                     "Couldn't map"}

    # Audit trail records the run.
    audit = wb["Audit Trail"]
    assert _table_rows(audit, _header_row(audit, "Seq"))


def test_health_report_is_formula_safe_and_logged(boundary_run):
    api, rid, r, _ = boundary_run
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    for name in ("Exceptions", "Duplicates", "Coverage"):
        for row in wb[name].iter_rows():
            for c in row:
                if c.data_type == "f":
                    raise AssertionError(f"unexpected formula in {name}!{c.coordinate}: {c.value}")
    trail = api.get(f"/api/v1/reports/{rid}/audit", params={"action_type": "EXPORT_GENERATED"})
    assert trail.status_code == 200, trail.text
    assert any((e.get("after_value") or e.get("after") or {}).get("export") == "health_report_xlsx"
               for e in trail.json()["items"]), trail.text


def _dup_book() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Claims"
    ws.append(["Claim Reference", "Insured Name", "Policy Number", "Date of Loss", "Claim Status", "Currency",
               "Paid to Date", "Reserve", "Total Incurred"])
    for i in range(10):
        ws.append([f"CLM-{i:04d}", f"Insured {i}", f"POL-{500 + i}", "2024-01-15", "Open", "GBP", 1000, 500, 1500])
    ws.append(["CLM-0003", "Insured 3", "POL-503", "2024-01-15", "Open", "GBP", 1000, 500, 1500])  # exact copy
    ws.append(["CLM-0103", "insured 7", "POL-507", "2024-01-16", "Open", "GBP", 1000, 500, 1500])  # probable
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_duplicates_sheet_groups(api: Api):
    rid, _ = api.full_run("dups.xlsx", _dup_book())
    summary = api.get(f"/api/v1/reports/{rid}/summary").json()["summary"]
    r = api.get(f"/api/v1/reports/{rid}/export/health-report.xlsx")
    assert r.status_code == 200, r.text
    rows = _table_rows(openpyxl.load_workbook(io.BytesIO(r.content))["Duplicates"], 1)
    pairs = int(summary["exact_duplicates"]) + int(summary["probable_duplicates"] or 0)
    assert pairs >= 1
    groups: dict[str, list] = {}
    for x in rows:
        groups.setdefault(x[0], []).append(x)
    assert len(groups) == pairs
    assert all(len(g) >= 2 and g[0][12] == "Keep" for g in groups.values())
    assert any(x[1] == "Exact" for x in rows)


def test_health_report_needs_a_processed_report_and_owner(api: Api, api_b: Api):
    rid, _ = api.full_run(BOUNDARY.name, BOUNDARY.read_bytes())
    assert api_b.get(f"/api/v1/reports/{rid}/export/health-report.xlsx").status_code == 404


@pytest.mark.slow
def test_health_report_10k_rows_under_5_seconds():
    from perf_health_report import synthetic_data

    from app.exports.health_report import build_health_report

    data = synthetic_data(rows=10_000, sheets=10)
    took = float("inf")
    for _ in range(3):  # best of three: the builder's speed, not a busy machine's
        t0 = time.perf_counter()
        body = build_health_report(data)
        took = min(took, time.perf_counter() - t0)
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True)
    assert wb.sheetnames == list(SHEETS)
    print(f"\n10,000-row Health Report built in {took:.2f}s ({len(body) / 1e6:.1f} MB)")
    assert took < 5


def test_annotated_csv_is_converted_and_says_so(api: Api):
    csv_text = ("Claim Reference,Insured Name,Policy Number,Date of Loss,Claim Status,Currency,Paid to Date,"
                "Reserve,Total Incurred\n"
                + "".join(f"CLM-{i:04d},Insured {i},{100000 + i},2024-01-15,Open,GBP,1000,500,1500\n"
                          for i in range(8)))
    r = api.upload("claims.csv", csv_text.encode(), "text/csv")
    assert r.status_code == 202, r.text
    rid = r.json()["id"]
    from conftest import run_jobs
    run_jobs()
    api.confirm_all(rid)
    assert api.post(f"/api/v1/reports/{rid}/process").status_code in (200, 202)
    run_jobs()
    a = api.get(f"/api/v1/reports/{rid}/export/annotated.xlsx")
    assert a.status_code == 200, a.text
    notes = openpyxl.load_workbook(io.BytesIO(a.content))["TrueBind Notes"]
    assert any(isinstance(c.value, str) and c.value.startswith("Your file was a .csv file")
               for row in notes.iter_rows() for c in row)
